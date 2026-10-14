"""外部情報的四種輸出：每日摘要 .md、每週彙整 .docx、情報追蹤表 .xlsx、summary.json。"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from intel.settings import GROUPS
from intel.store import DayStore
from shared.report import Footer, Section, write_docx, write_md, write_xlsx

LEVEL = "公開"
TRACK_SHEET = "追蹤表"
UNREAD = "未讀"
TRACK_COLUMNS = [
    "文章編號",
    "首次看到",
    "發布日期",
    "來源",
    "標題",
    "連結",
    "分類",
    "是否相關",
    "摘要",
    "對優樂地的意義",
    "分數",
    "命中關鍵字",
    "已讀",
]
MOCK_NOTE = "注意：正式摘要模型尚未設定，本期摘要為「模擬摘要」，只截取原文第一句，不是正式內容。"


# ---- 共用 ----
def _method(th: dict) -> list[str]:
    ms = th.get("relevance_min_score")
    sim = th.get("title_similarity_dup")
    cap = th.get("max_items_per_issue")
    return [
        "相關度＝命中關鍵字的權重加總（同一關鍵字出現多次只算一次；同義詞算同一關鍵字）。",
        f"相關門檻：{'未設定，全部列為相關' if ms is None else f'{ms} 分以上'}（待主管確認）。",
        f"同一事件：網址去掉追蹤參數後相同，或標題相似度 {'未設定（不比對）' if sim is None else f'≥ {sim}'}（待主管確認）。",
        "分組：依命中權重最高的關鍵字類別；沒有命中時用來源的分類。",
        f"每期上限：{'不設上限' if cap is None else f'{cap} 則，超出者列入附錄'}（待主管確認）。",
        "所有連結都直接取自來源的 RSS，不由模型產生。",
    ]


def _missing(day: dict) -> list[str]:
    out = []
    if day.get("mock"):
        out.append("正式摘要模型尚未設定（本期為模擬摘要）")
    th = day.get("thresholds", {})
    for k, label in (
        ("relevance_min_score", "相關度門檻"),
        ("title_similarity_dup", "標題相似度門檻"),
        ("max_items_per_issue", "每期最多則數"),
    ):
        if th.get(k) is None:
            out.append(f"{label}未設定")
    out += [f"設定問題：{p}" for p in day.get("config_problems", [])]
    for r in day.get("records", []):
        if r["relevant"] and r["summary"] and r["summary"]["status"] != "ok":
            out.append(f"「{r['title']}」：{r['summary']['summary']}")
    return out


def _footer(day: dict, as_of: dt.date) -> Footer:
    return Footer(
        as_of=as_of,
        input_files=day.get("input_files", []),
        run_id=day["run_id"],
        level=LEVEL,
        missing=_missing(day),
        method=_method(day.get("thresholds", {})),
    )


def _anomaly_lines(day: dict) -> list[str]:
    lines = [f"- {a['source']}：{a['reason']}" for a in day.get("anomalies", [])]
    lines += [f"- 設定問題：{p}" for p in day.get("config_problems", [])]
    return lines or ["無"]


def _item_lines(r: dict, md: bool) -> list[str]:
    title = f"**[{r['title']}]({r['url']})**" if md else f"■ {r['title']}"
    terms = "、".join(r["matched_terms"]) or "無"
    lines = [
        title,
        f"來源：{'、'.join(r['sources'])}｜發布日期：{r['published'] or '待補'}｜相關度 {r['score']} 分（命中：{terms}）",
    ]
    s = r.get("summary") or {}
    lines.append(f"摘要：{s.get('summary', '待補')}")
    lines.append(f"對優樂地的意義：{s.get('meaning', '待補')}")
    if s.get("flags"):
        lines.append(f"⚠ {'；'.join(s['flags'])}")
    if not md:
        lines.append(f"原文連結：{r['url']}")
    if r["extra_urls"]:
        lines.append("其他來源連結：" + "、".join(r["extra_urls"]))
    return lines


def _group_sections(records: list[dict], md: bool) -> list[Section]:
    rel = [r for r in records if r["relevant"]]
    if not rel:
        return [Section("相關動態", ["本期無符合項目。"])]
    out = []
    for g in GROUPS:
        rs = [r for r in rel if r["group"] == g]
        if not rs:
            continue
        paras: list[str] = []
        for r in rs:
            paras += _item_lines(r, md)
        out.append(Section(f"{g}（{len(rs)} 則）", paras))
    return out


def _low_section(records: list[dict]) -> Section:
    low = [r for r in records if not r["relevant"]]
    if not low:
        return Section("附錄：低相關與未列入", ["無"])
    df = pd.DataFrame(
        [
            {
                "標題": r["title"],
                "來源": "、".join(r["sources"]),
                "分數": r["score"],
                "原因": r["excluded_reason"],
                "連結": r["url"],
            }
            for r in low
        ]
    )
    return Section("附錄：低相關與未列入", ["以下文章沒有放進本期摘要，但保留供查閱。"], df)


def _headline(records: list[dict], anomalies: list, period: str) -> str:
    rel = [r for r in records if r["relevant"]]
    if not rel:
        head = f"{period}沒有符合的外部情報（本期無符合項目）"
    else:
        parts = [f"{g} {n} 則" for g in GROUPS if (n := sum(r["group"] == g for r in rel))]
        head = f"{period}有 {len(rel)} 則相關動態：{'、'.join(parts)}"
    errs = sum(1 for a in anomalies if a.get("level") == "error")
    if errs:
        head += f"；另有 {errs} 個來源抓取失敗，請見「來源異常」"
    return head + "。"


# ---- 每日 ----
def write_daily(day: dict, out_dir: Path, as_of: dt.date) -> Path:
    records = day["records"]
    sections = []
    if day.get("mock"):
        sections.append(Section("注意", [MOCK_NOTE]))
    sections.append(Section("來源異常", _anomaly_lines(day)))
    sections += _group_sections(records, md=True)
    sections.append(_low_section(records))
    return write_md(
        out_dir / f"每日摘要_{as_of.strftime('%Y%m%d')}.md",
        f"外部情報每日摘要（{as_of.isoformat()}）",
        _headline(records, day.get("anomalies", []), "今天"),
        sections,
        _footer(day, as_of),
    )


# ---- 每週 ----
def _week(days: DayStore, as_of: dt.date) -> tuple[list[dict], list[dict], dict]:
    window = days.window(as_of, 7)
    records = [r for d in window for r in d["records"]]
    records.sort(key=lambda r: (not r["relevant"], -r["score"], r["published"] or "", r["id"]))
    anomalies = [
        {**a, "source": f"{d['as_of']} {a['source']}"}
        for d in window
        for a in d.get("anomalies", [])
    ]
    return records, anomalies, window[-1] if window else {}


def write_weekly(day: dict, days: DayStore, out_dir: Path, as_of: dt.date) -> Path:
    records, anomalies, _ = _week(days, as_of)
    start = as_of - dt.timedelta(days=6)
    wk = {**day, "records": records, "anomalies": anomalies}
    sections = []
    if any(d.get("mock") for d in days.window(as_of, 7)):
        sections.append(Section("注意", [MOCK_NOTE]))
    sections.append(Section("來源異常", _anomaly_lines(wk)))
    sections += _group_sections(records, md=False)
    sections.append(_low_section(records))
    return write_docx(
        out_dir / f"每週彙整_{as_of.strftime('%Y%m%d')}.docx",
        f"外部情報每週彙整（{start.isoformat()}～{as_of.isoformat()}）",
        _headline(records, anomalies, "本週"),
        sections,
        _footer(wk, as_of),
    )


# ---- 追蹤表 ----
def _read_marks(path: Path) -> dict[str, str]:
    """讀回主管在追蹤表填的「已讀」，依文章編號對回。"""
    if not path.exists():
        return {}
    try:
        ws = load_workbook(path, read_only=True)[TRACK_SHEET]
    except (KeyError, OSError, ValueError):
        return {}
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return {}
    header = [str(c) if c is not None else "" for c in rows[1]]
    if "文章編號" not in header or "已讀" not in header:
        return {}
    i_id, i_read = header.index("文章編號"), header.index("已讀")
    marks = {}
    for row in rows[2:]:
        if not row or row[i_id] in (None, ""):
            break
        v = row[i_read]
        if v not in (None, "", "待補", UNREAD):
            marks[str(row[i_id])] = str(v)
    return marks


def write_tracking(day: dict, days: DayStore, out_dir: Path, as_of: dt.date) -> Path:
    path = out_dir / "情報追蹤表.xlsx"
    marks = _read_marks(path)
    rows = []
    for d in days.all_until(as_of):
        for r in d["records"]:
            s = r.get("summary") or {}
            rows.append(
                {
                    "文章編號": r["id"],
                    "首次看到": r["first_seen"],
                    "發布日期": r["published"],
                    "來源": "、".join(r["sources"]),
                    "標題": r["title"],
                    "連結": r["url"],
                    "分類": r["group"],
                    "是否相關": "相關" if r["relevant"] else "低相關",
                    "摘要": s.get("summary"),
                    "對優樂地的意義": s.get("meaning"),
                    "分數": r["score"],
                    "命中關鍵字": "、".join(r["matched_terms"]) or "無",
                    "已讀": marks.get(r["id"], UNREAD),
                }
            )
    # 最新的在最上面；同一天內相關的在前、分數高的在前
    rows.sort(key=lambda x: (x["是否相關"] != "相關", -x["分數"], x["文章編號"]))
    rows.sort(key=lambda x: x["首次看到"], reverse=True)
    df = pd.DataFrame(rows, columns=TRACK_COLUMNS)
    n_rel = sum(1 for r in rows if r["是否相關"] == "相關")
    headline = f"追蹤表共 {len(rows)} 則，其中相關 {n_rel} 則。看過的請把「已讀」欄的「未讀」改成「已讀」，重新產生時會保留。"
    return write_xlsx(path, headline, {TRACK_SHEET: df}, _footer(day, as_of))


# ---- summary.json（給管理層週報）----
def _window_counts(days: DayStore, end: dt.date) -> tuple[int, int, bool]:
    w = days.window(end, 7)
    rel = [r for d in w for r in d["records"] if r["relevant"]]
    return len(rel), sum(r["group"] == "法規與準則" for r in rel), bool(w)


def write_summary(day: dict, days: DayStore, out_dir: Path, as_of: dt.date) -> Path:
    start = as_of - dt.timedelta(days=6)
    cur_rel, cur_reg, _ = _window_counts(days, as_of)
    prev_rel, prev_reg, has_prev = _window_counts(days, as_of - dt.timedelta(days=7))
    records, anomalies, _ = _week(days, as_of)
    ref = f"data/interim/intel/days/ {start.isoformat()}～{as_of.isoformat()}"
    data = {
        "system": "intel",
        "run_id": day["run_id"],
        "as_of": as_of.isoformat(),
        "classification": LEVEL,
        "headline": _headline(records, anomalies, "近 7 天"),
        "metrics": [
            {
                "key": "intel.relevant_items",
                "label": "近 7 天相關外部情報則數",
                "value": cur_rel,
                "unit": "則",
                "prev_value": prev_rel if has_prev else None,
                "source_ref": ref,
            },
            {
                "key": "intel.regulation_items",
                "label": "其中法規與準則類則數",
                "value": cur_reg,
                "unit": "則",
                "prev_value": prev_reg if has_prev else None,
                "source_ref": ref,
            },
        ],
        "alerts": [],
        "data_quality": [f"{a['source']}：{a['reason']}" for a in anomalies]
        + [f"設定問題：{p}" for p in day.get("config_problems", [])],
        "highlights": [
            {"title": r["title"], "url": r["url"], "group": r["group"]}
            for r in records
            if r["relevant"]
        ][:3],
        "mock_summary": bool(day.get("mock")),
    }
    path = out_dir / "summary.json"
    out_dir.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def write_all(
    day: dict, days: DayStore, out_dir: Path, as_of: dt.date, weekly: bool
) -> dict[str, Path]:
    files = {
        "daily": write_daily(day, out_dir, as_of),
        "tracking": write_tracking(day, days, out_dir, as_of),
        "summary": write_summary(day, days, out_dir, as_of),
    }
    if weekly:
        files["weekly"] = write_weekly(day, days, out_dir, as_of)
    return files
