"""共用輸出版型：.xlsx、.docx、.md。

每份輸出都帶頁尾：資料截至日期、輸入檔名、執行編號、資料等級、待補清單，
可另附「計算方法」段。表格中的空值一律顯示「待補」。
"""

from __future__ import annotations

import datetime as dt
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from shared.classify import Classifier
from shared.dates import MISSING, fmt_date

LEVEL_BANNER = {
    "公開": "資料等級：公開",
    "內部一般": "資料等級：內部一般",
    "內部敏感": "內部敏感：僅執行長與主管",
    "客戶機密": "客戶機密：僅執行長與主管",
}
_SHEET_NAME_MAX = 31


def make_run_id(system: str, as_of: dt.date, input_files: list[Path]) -> str:
    """執行編號＝系統-日期-輸入檔內容雜湊前 8 碼。同樣輸入得到同樣編號。"""
    h = hashlib.sha256()
    for p in sorted(Path(x) for x in input_files):
        h.update(p.name.encode("utf-8"))
        h.update(Path(p).read_bytes() if Path(p).exists() else b"<missing>")
    return f"{system}-{as_of.strftime('%Y%m%d')}-{h.hexdigest()[:8]}"


@dataclass
class Footer:
    as_of: dt.date
    input_files: list[str]
    run_id: str
    level: str
    missing: list[str] = field(default_factory=list)  # 待補清單
    method: list[str] = field(default_factory=list)  # 計算方法

    def __post_init__(self) -> None:
        # 不認得或沒給的等級，一律當作客戶機密標示
        self.level = Classifier().normalize(self.level)

    def lines(self) -> list[str]:
        out = [
            f"資料截至：{fmt_date(self.as_of)}",
            f"輸入檔：{'、'.join(self.input_files) if self.input_files else '無'}",
            f"執行編號：{self.run_id}",
            f"資料等級：{self.level}",
            "待補清單：" + ("無" if not self.missing else ""),
        ]
        out += [f"  - {m}" for m in self.missing]
        if self.method:
            out.append("計算方法：")
            out += [f"  - {m}" for m in self.method]
        return out


@dataclass
class Section:
    title: str
    paragraphs: list[str] = field(default_factory=list)
    table: pd.DataFrame | None = None


def fill_missing(df: pd.DataFrame) -> pd.DataFrame:
    """空值、空字串一律顯示「待補」；日期轉成 YYYY-MM-DD。"""
    out = df.copy().astype(object)
    for col in out.columns:
        out[col] = [_cell(v) for v in out[col]]
    return out


def _cell(v: object) -> object:
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return MISSING
    if isinstance(v, str) and not v.strip():
        return MISSING
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return v


def write_xlsx(path: Path, headline: str, sheets: dict[str, pd.DataFrame], footer: Footer) -> Path:
    """第一個分頁放等級標籤、結論、頁尾；之後每個分頁一張表，最後附頁尾。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "總覽"
    ws.append([LEVEL_BANNER.get(footer.level, footer.level)])
    ws["A1"].font = Font(bold=True, color="C00000")
    ws.append([headline])
    ws["A2"].font = Font(bold=True, size=13)
    ws.append([])
    for line in footer.lines():
        ws.append([line])
    for name, df in sheets.items():
        sh = wb.create_sheet(name[:_SHEET_NAME_MAX])
        sh.append([LEVEL_BANNER.get(footer.level, footer.level)])
        df = fill_missing(df)
        sh.append([str(c) for c in df.columns])
        for c in range(1, len(df.columns) + 1):
            sh.cell(row=2, column=c).font = Font(bold=True)
            sh.cell(row=2, column=c).fill = PatternFill("solid", fgColor="DDEBF7")
        for row in df.itertuples(index=False):
            sh.append(list(row))
        sh.append([])
        for line in footer.lines():
            sh.append([line])
        for i, col in enumerate(df.columns, start=1):
            sh.column_dimensions[get_column_letter(i)].width = max(
                10, min(40, len(str(col)) * 2 + 4)
            )
    wb.properties.creator = "unity-mgmt-suite"
    wb.properties.created = dt.datetime(2000, 1, 1)  # 固定時間，讓重跑結果可比對
    wb.properties.modified = dt.datetime(2000, 1, 1)
    wb.save(path)
    return path


def write_docx(
    path: Path, title: str, headline: str, sections: list[Section], footer: Footer
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = LEVEL_BANNER.get(footer.level, footer.level)
    doc.add_heading(title, level=1)
    doc.add_paragraph(headline).runs[0].bold = True
    for s in sections:
        doc.add_heading(s.title, level=2)
        for p in s.paragraphs:
            doc.add_paragraph(p)
        if s.table is not None:
            df = fill_missing(s.table)
            t = doc.add_table(rows=1, cols=len(df.columns))
            t.style = "Table Grid"
            for i, c in enumerate(df.columns):
                t.rows[0].cells[i].text = str(c)
            for row in df.itertuples(index=False):
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
    doc.add_heading("頁尾資訊", level=3)
    for line in footer.lines():
        doc.add_paragraph(line)
    doc.core_properties.author = "unity-mgmt-suite"
    doc.core_properties.created = dt.datetime(2000, 1, 1)
    doc.core_properties.modified = dt.datetime(2000, 1, 1)
    doc.save(path)
    return path


def to_markdown(title: str, headline: str, sections: list[Section], footer: Footer) -> str:
    out = [
        f"> {LEVEL_BANNER.get(footer.level, footer.level)}",
        "",
        f"# {title}",
        "",
        f"**{headline}**",
        "",
    ]
    for s in sections:
        out += [f"## {s.title}", ""]
        out += [p + "\n" for p in s.paragraphs]
        if s.table is not None:
            df = fill_missing(s.table)
            out.append("| " + " | ".join(str(c) for c in df.columns) + " |")
            out.append("|" + "---|" * len(df.columns))
            for row in df.itertuples(index=False):
                out.append("| " + " | ".join(str(v).replace("|", "／") for v in row) + " |")
            out.append("")
    out += ["---", ""] + [f"{line}  " for line in footer.lines()]
    return "\n".join(out) + "\n"


def write_md(
    path: Path, title: str, headline: str, sections: list[Section], footer: Footer
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(to_markdown(title, headline, sections, footer), encoding="utf-8")
    return path
