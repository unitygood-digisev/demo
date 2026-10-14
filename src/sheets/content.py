"""從本機已產生的報表組出各分頁內容（不重新計算）。"""

from __future__ import annotations

import datetime as dt
import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

MISSING = "待補"
INTEL_COLUMNS = [
    "首次看到",
    "來源",
    "標題",
    "連結",
    "分類",
    "是否相關",
    "摘要",
    "對優樂地的意義",
    "分數",
]


class SyncError(Exception):
    """同步無法進行（訊息為白話中文）。"""


@dataclass
class Tab:
    title: str
    rows: list[list]  # 第 1 列：等級與資料截至；之後是表格


def _cell(v: object) -> object:
    if v is None or (isinstance(v, str) and not v.strip()):
        return MISSING
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()[:10]
    return v


def read_sheet_table(path: Path, sheet: str) -> tuple[list[str], list[list]]:
    """讀本系統產生的 xlsx 分頁：第 2 列是表頭，資料到第一個整列空白為止。"""
    try:
        ws = load_workbook(path, read_only=True)[sheet]
    except (OSError, KeyError, ValueError) as e:
        raise SyncError(f"無法讀取 {path.name} 的「{sheet}」分頁") from e
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return [], []
    header = [str(c) if c is not None else "" for c in rows[1]]
    data = []
    for r in rows[2:]:
        if not r or all(c is None for c in r):
            break
        data.append([_cell(c) for c in r[: len(header)]])
    return header, data


def latest(folder: Path, pattern: str) -> Path | None:
    files = sorted(folder.glob(pattern)) if folder.exists() else []
    return files[-1] if files else None


def load_digest(digest_dir: Path, as_of: dt.date) -> dict:
    p = digest_dir / f"管理層週報_資料_{as_of.strftime('%Y%m%d')}.json"
    if not p.exists():
        raise SyncError(f"找不到 {as_of.isoformat()} 的週報資料檔，請先產生週報再同步")
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        raise SyncError(f"週報資料檔 {p.name} 無法讀取") from e


def _banner(d: dict) -> list:
    return [f"{d['level_banner']}｜資料截至 {d['as_of']}｜執行編號 {d['run_id']}"]


def build_tabs(
    names: dict[str, str], digest: dict, progress_xlsx: Path | None, intel_xlsx: Path | None
) -> list[Tab]:
    top = _banner(digest)
    tabs: list[Tab] = []

    rows: list[list] = [top, [digest["title"]], ["本週結論"]]
    rows += [[f"{i}. {s}"] for i, s in enumerate(digest["sentences"], 1)]
    rows += [[], ["關鍵指標", "本週", "上週", "變化", "單位"]]
    rows += [
        [m["label"], m["this"], m["prev"], m["change"], m["unit"] or "—"] for m in digest["metrics"]
    ]
    rows += [[], ["需要您決定的事項"]] + (
        [[x] for x in digest["decisions"]] or [["本週沒有紅色預警。"]]
    )
    if digest["decisions_more"]:
        rows.append([f"（其餘 {digest['decisions_more']} 條紅色預警，見「進度預警」分頁）"])
    rows += [[], ["外部情報", "連結"]] + [[i["title"], i["url"]] for i in digest["intel"]]
    if digest.get("intel_note"):
        rows.append([digest["intel_note"]])
    rows += [[], ["資料品質提醒"]] + (
        [[x] for x in digest["data_quality"]] or [["沒有資料品質提醒。"]]
    )
    if digest["data_quality_more"]:
        rows.append([f"（其餘 {digest['data_quality_more']} 項，見各系統報表）"])
    tabs.append(Tab(names["digest"], rows))

    for key, sheet, label in (
        ("alerts", "預警清單", "進度預警"),
        ("projects", "項目明細", "進度明細"),
    ):
        if progress_xlsx is None:
            tabs.append(Tab(names[key], [top, [f"項目進度看板本週沒有產生報表，{label}未更新"]]))
            continue
        header, data = read_sheet_table(progress_xlsx, sheet)
        tabs.append(Tab(names[key], [top, [f"來源：{progress_xlsx.name}"], header, *data]))

    if intel_xlsx is None:
        tabs.append(Tab(names["intel"], [top, ["外部情報沒有產生追蹤表，未更新"]]))
    else:
        header, data = read_sheet_table(intel_xlsx, "追蹤表")
        idx = [header.index(c) for c in INTEL_COLUMNS if c in header]
        tabs.append(
            Tab(
                names["intel"],
                [
                    top,
                    [f"來源：{intel_xlsx.name}"],
                    [header[i] for i in idx],
                    *[[r[i] for i in idx] for r in data],
                ],
            )
        )

    info = [
        top,
        [
            "本試算表由管理層系統每週自動寫入（單向）。請勿在系統分頁中編輯，下次同步會被覆蓋；可另外新增自己的分頁。"
        ],
        [],
    ]
    info += [[line] for line in digest["footer"]]
    tabs.append(Tab(names["info"], info))
    return tabs


def history_row(digest: dict) -> tuple[list[str], list]:
    """指標歷史：每週一列。"""
    header = ["週報日期"] + [m["label"] for m in digest["metrics"]] + ["資料等級"]
    row = [digest["as_of"]] + [m["this"] for m in digest["metrics"]] + [digest["level"]]
    return header, row


def merge_history(existing: list[list], header: list[str], row: list) -> list[list]:
    """把本週那一列併入既有的指標歷史：同一週覆蓋、不同週新增；欄位變動時保留舊欄。"""
    old_header = [str(c) for c in existing[0]] if existing else []
    cols = old_header + [c for c in header if c not in old_header] if old_header else list(header)
    records: dict[str, dict] = {}
    for r in existing[1:]:
        if r and str(r[0]).strip():
            records[str(r[0])] = {old_header[i]: r[i] for i in range(min(len(r), len(old_header)))}
    records[str(row[0])] = {header[i]: row[i] for i in range(len(header))}
    out = [cols]
    for k in sorted(records):
        out.append([records[k].get(c, "") for c in cols])
    return out
