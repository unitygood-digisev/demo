"""讀取公司現有的項目進度看板（只讀，不改原檔）。"""

from __future__ import annotations

import datetime as dt
import fnmatch
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from shared.config import load_yaml
from shared.dates import DateParseError, parse_date
from shared.io import ColumnSpec, TableSpec, ValidationIssue, load_table


class ProgressDataError(Exception):
    """需要人處理才能繼續的問題（訊息為白話中文）。"""


@dataclass(frozen=True)
class BoardConfig:
    sheet: str
    input_glob: str
    header_anchor: str
    columns: dict[str, str]
    excluded_columns: tuple[str, ...]
    stages: tuple[str, ...]
    done_stages: tuple[str, ...]
    blocked_stages: tuple[str, ...]
    active_stages: tuple[str, ...]
    non_person_owners: tuple[str, ...]
    board_date_pattern: str

    @classmethod
    def load(cls, config_dir: Path | None = None) -> BoardConfig:
        c = load_yaml("progress_columns.yaml", config_dir)
        return cls(
            sheet=c["sheet"],
            input_glob=c["input_glob"],
            header_anchor=c["header_anchor"],
            columns=dict(c["columns"]),
            excluded_columns=tuple(c.get("excluded_columns") or ()),
            stages=tuple(c["stages"]),
            done_stages=tuple(c["done_stages"]),
            blocked_stages=tuple(c["blocked_stages"]),
            active_stages=tuple(c["active_stages"]),
            non_person_owners=tuple(c.get("non_person_owners") or ()),
            board_date_pattern=c["board_date_pattern"],
        )

    def spec(self) -> TableSpec:
        return TableSpec(
            "項目看板",
            (
                ColumnSpec("案子"),
                ColumnSpec("類別", required=False, must_exist=True),
                ColumnSpec("階段", "enum", allowed=self.stages),
                ColumnSpec("負責人", required=False, must_exist=True),
                ColumnSpec("下一步", required=False, must_exist=True),
                ColumnSpec("下一個日期", "date", required=False, must_exist=True),
            ),
            unique_key="案子",
        )


@dataclass
class Board:
    path: Path
    rows: list[dict]  # 每列：_row、案子、類別、階段、負責人、下一步、下一個日期
    issues: list[ValidationIssue] = field(default_factory=list)
    board_date: dt.date | None = None
    duplicates: set[str] = field(default_factory=set)


def find_input(raw_dir: Path, cfg: BoardConfig) -> tuple[Path, list[str]]:
    """回傳 (要讀的看板, 提醒)。只讀 .xlsx（副檔名不分大小寫）；同名開頭的其他格式檔案會提醒被略過。"""
    stem_glob = cfg.input_glob.rsplit(".", 1)[0].lower() + ".*"
    candidates = [
        p
        for p in (raw_dir.iterdir() if raw_dir.exists() else [])
        if p.is_file()
        and not p.name.startswith("~$")
        and fnmatch.fnmatch(p.name.lower(), stem_glob)
    ]
    files = sorted(p for p in candidates if fnmatch.fnmatch(p.name.lower(), cfg.input_glob.lower()))
    others = sorted(p.name for p in candidates if p not in files)
    notes = (
        [f"data/raw 裡有不是 .xlsx 的看板檔（{'、'.join(others)}），系統只讀 .xlsx，已略過"]
        if others
        else []
    )
    if not files:
        raise ProgressDataError(
            f"在 {raw_dir} 找不到符合「{cfg.input_glob}」的看板檔案，請先把看板放進去"
        )
    return files[-1], notes


def _board_date(path: Path, cfg: BoardConfig) -> dt.date | None:
    """從標題列上方的說明文字找看板資料日期，例如「來源：2026/10/5 會議紀錄」。"""
    try:
        wb = load_workbook(path, read_only=True)
        ws = wb[cfg.sheet]
    except Exception:
        return None
    pat = re.compile(cfg.board_date_pattern)
    texts: list[str] = []
    for row in ws.iter_rows(max_row=100, values_only=True):  # 掃到標題列為止
        cells = [str(c).strip() for c in row if c is not None]
        if cfg.header_anchor in cells and len(set(cells) & set(cfg.columns)) >= 2:
            break  # 到標題列為止（說明文字剛好寫了錨點的那一列不算）
        texts.extend(cells)
    # 優先取寫有「來源」的那一格，其次才是其他說明文字裡的第一個日期
    for text in sorted(texts, key=lambda t: "來源" not in t):
        for m in pat.finditer(text):
            try:
                return parse_date(m.group(0))
            except DateParseError:
                continue
    return None


def _clean(v):
    """表格轉回 Python 值：空值一律 None，時間戳轉日期。"""
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return v.date()
    return v


def load_board(path: Path, cfg: BoardConfig) -> Board:
    res = load_table(
        path,
        cfg.spec(),
        cfg.columns,
        sheet=cfg.sheet,
        header_anchor=cfg.header_anchor,
        header_expect=tuple(cfg.columns),
    )
    rows = [{k: _clean(v) for k, v in r.items()} for r in res.data.to_dict("records")]
    rows = [r for r in rows if r.get("案子")]
    names = [r["案子"] for r in rows]
    dups = {n for n in names if names.count(n) > 1}
    return Board(path, rows, list(res.issues), _board_date(path, cfg), dups)
