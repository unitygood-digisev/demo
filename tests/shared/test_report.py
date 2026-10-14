import datetime as dt

import pandas as pd
from docx import Document
from openpyxl import load_workbook

from shared.report import (
    Footer,
    Section,
    make_run_id,
    to_markdown,
    write_docx,
    write_md,
    write_xlsx,
)

FOOTER = Footer(
    as_of=dt.date(2026, 10, 14),
    input_files=["projects.xlsx", "milestones.xlsx"],
    run_id="progress-20261014-abcd1234",
    level="內部一般",
    missing=["P010 負責人"],
    method=["逾期 = 預計日 < 今天 且 狀態 ≠ 已完成"],
)
DF = pd.DataFrame(
    {"項目": ["P1", "P2"], "負責人": ["王", None], "日期": [dt.date(2026, 1, 2), None]}
)


def _all_text_xlsx(path):
    wb = load_workbook(path)
    return {
        ws.title: [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        for ws in wb
    }


def test_xlsx_footer_and_missing(tmp_path):
    p = write_xlsx(tmp_path / "out.xlsx", "本週有 1 個項目可能逾期", {"明細": DF}, FOOTER)
    sheets = _all_text_xlsx(p)
    assert set(sheets) == {"總覽", "明細"}
    for cells in sheets.values():
        joined = "\n".join(map(str, cells))
        for needed in [
            "資料截至：2026-10-14",
            "projects.xlsx",
            "progress-20261014-abcd1234",
            "資料等級：內部一般",
            "P010 負責人",
        ]:
            assert needed in joined
    assert sheets["總覽"][1] == "本週有 1 個項目可能逾期"  # 第一句講結論
    assert sheets["明細"].count("待補") == 2
    assert "2026-01-02" in sheets["明細"]
    assert "逾期 = 預計日 < 今天 且 狀態 ≠ 已完成" in "\n".join(map(str, sheets["總覽"]))


def test_sensitive_banner(tmp_path):
    f = Footer(dt.date(2026, 10, 14), ["a.xlsx"], "x", "客戶機密")
    p = write_xlsx(tmp_path / "o.xlsx", "結論", {}, f)
    assert load_workbook(p)["總覽"]["A1"].value == "客戶機密：僅執行長與主管"


def test_docx(tmp_path):
    p = write_docx(
        tmp_path / "o.docx", "標題", "結論句", [Section("明細", ["說明文字"], DF)], FOOTER
    )
    doc = Document(p)
    assert doc.sections[0].header.paragraphs[0].text == "資料等級：內部一般"
    texts = [para.text for para in doc.paragraphs]
    assert texts[1] == "結論句"
    assert "執行編號：progress-20261014-abcd1234" in texts
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert cells.count("待補") == 2


def test_markdown(tmp_path):
    md = to_markdown("標題", "結論句", [Section("明細", [], DF)], FOOTER)
    assert md.splitlines()[0] == "> 資料等級：內部一般"
    assert "**結論句**" in md and "待補" in md and "待補清單：" in md
    p = write_md(tmp_path / "o.md", "標題", "結論句", [], FOOTER)
    assert p.read_text(encoding="utf-8").startswith("> 資料等級")


def test_footer_no_missing_says_none():
    f = Footer(dt.date(2026, 10, 14), [], "r", "公開")
    assert "待補清單：無" in f.lines()


def test_run_id_reproducible(tmp_path):
    a = tmp_path / "a.xlsx"
    a.write_bytes(b"123")
    r1 = make_run_id("progress", dt.date(2026, 10, 14), [a])
    r2 = make_run_id("progress", dt.date(2026, 10, 14), [a])
    assert r1 == r2 and r1.startswith("progress-20261014-")
    a.write_bytes(b"124")
    assert make_run_id("progress", dt.date(2026, 10, 14), [a]) != r1
