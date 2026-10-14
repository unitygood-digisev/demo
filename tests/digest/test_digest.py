"""系統 5 管理層週報測試。"""

import datetime as dt
import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
import yaml
from docx import Document

from digest.build import fmt_change
from digest.pipeline import RunConfig, run

ROOT = Path(__file__).resolve().parents[2]
AS_OF = dt.date(2026, 10, 14)


def summary(system, level="內部一般", as_of=AS_OF, **kw):
    base = {
        "system": system,
        "run_id": f"{system}-x",
        "as_of": as_of.isoformat(),
        "classification": level,
        "headline": f"{system} 的結論句。",
        "metrics": [],
        "alerts": [],
        "data_quality": [],
    }
    base.update(kw)
    return base


PROGRESS = summary(
    "progress",
    metrics=[
        {
            "key": "progress.red_alerts",
            "label": "紅",
            "value": 11,
            "unit": "案",
            "prev_value": 9,
            "source_ref": "x",
        },
        {
            "key": "progress.yellow_alerts",
            "label": "黃",
            "value": 4,
            "unit": "案",
            "prev_value": None,
            "source_ref": "x",
        },
        {
            "key": "progress.active_projects",
            "label": "進行中",
            "value": 32,
            "unit": "案",
            "prev_value": 32,
            "source_ref": "x",
        },
    ],
    alerts=[
        {
            "level": "紅",
            "rule": "P01",
            "why": f"原因{i}",
            "subject": f"案{i}",
            "suggested_owner": f"人{i}",
        }
        for i in range(7)
    ]
    + [
        {
            "level": "黃",
            "rule": "P02",
            "why": "黃色原因",
            "subject": "黃案",
            "suggested_owner": "某人",
        }
    ],
    data_quality=["看板資料日期是 2026-10-05，距今 9 天"],
)
INTEL = summary(
    "intel",
    level="公開",
    metrics=[
        {
            "key": "intel.relevant_items",
            "label": "情報",
            "value": 5,
            "unit": "則",
            "prev_value": 7,
            "source_ref": "x",
        }
    ],
    data_quality=["來源 A 抓取失敗"],
    highlights=[
        {"title": f"新聞{i}", "url": f"https://example.com/{i}", "group": "其他"} for i in range(5)
    ],
    mock_summary=True,
)


def write(path: Path, data) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        data if isinstance(data, str) else json.dumps(data, ensure_ascii=False), encoding="utf-8"
    )
    return path


@pytest.fixture
def env(tmp_path):
    cfg = tmp_path / "config"
    shutil.copytree(ROOT / "config", cfg)
    p = write(tmp_path / "progress.json", PROGRESS)
    i = write(tmp_path / "intel.json", INTEL)

    def make(**kw):
        overrides = {"progress": p, "intel": i}
        overrides.update(kw.pop("overrides", {}))
        return RunConfig(
            as_of=kw.pop("as_of", AS_OF),
            out_dir=tmp_path / "out",
            config_dir=cfg,
            summary_overrides=overrides,
            log_dir=tmp_path / "logs",
            **kw,
        )

    return {"make": make, "cfg": cfg, "tmp": tmp_path, "progress": p, "intel": i}


def docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs] + [
        c.text for t in d.tables for r in t.rows for c in r.cells
    ]
    return "\n".join(parts), d.sections[0].header.paragraphs[0].text


def test_every_number_equals_source(env):
    res = run(env["make"]())
    src = {m["key"]: m for s in (PROGRESS, INTEL) for m in s["metrics"]}
    rows = [m for m in res.digest.metrics if m.raw_value is not None]
    assert {m.key for m in rows} == {
        "progress.red_alerts",
        "progress.yellow_alerts",
        "progress.active_projects",
        "intel.relevant_items",
    }
    for m in rows:
        assert m.raw_value == src[m.key]["value"]
        assert m.this == str(src[m.key]["value"])
        prev = src[m.key]["prev_value"]
        assert m.prev == ("待補" if prev is None else str(prev))
    by = {m.key: m for m in rows}
    assert by["progress.red_alerts"].change == "+2"
    assert by["progress.yellow_alerts"].change == "待補"
    assert by["progress.active_projects"].change == "持平"
    assert by["intel.relevant_items"].change == "−2"
    text, _ = docx_text(res.files["docx"])
    txt = res.files["txt"].read_text(encoding="utf-8")
    assert "紅色預警項目數：11／9／+2（案）" in txt
    for v in ("11", "32", "−2"):
        assert v in text


def test_layout_sections_and_limits(env):
    res = run(env["make"]())
    d = res.digest
    assert d.sentences[0] == "項目進度看板：progress 的結論句。"
    assert d.sentences[1] == "外部情報摘要：intel 的結論句。"
    assert "項目財務總覽尚未上線" in d.sentences[2] and "2 項資料品質提醒" in d.sentences[2]
    assert len(d.decisions) == 5 and d.decisions_more == 2  # 只取紅色、前 5 條
    assert all("黃案" not in x for x in d.decisions)
    assert d.decisions[0] == "［項目進度看板］案0：原因0 → 建議跟 人0 確認"
    assert [i["title"] for i in d.intel] == ["新聞0", "新聞1", "新聞2"]
    assert [m.this for m in d.metrics if "財務" in m.label] == ["系統尚未上線"]
    assert sum(1 for m in d.metrics if "財務" in m.label) == 1  # 未上線只顯示一行
    assert any("模擬版" in x for x in d.missing)


def test_level_is_highest_input(env):
    res = run(env["make"]())
    assert res.digest.level == "內部一般"
    assert docx_text(res.files["docx"])[1] == "資料等級：內部一般"
    # 財務上線且是內部敏感
    cfg = yaml.safe_load((env["cfg"] / "digest.yaml").read_text(encoding="utf-8"))
    cfg["systems"][2]["launched"] = True
    (env["cfg"] / "digest.yaml").write_text(
        yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8"
    )
    f = write(env["tmp"] / "finance.json", summary("finance", level="內部敏感"))
    res = run(env["make"](overrides={"finance": f}))
    assert res.digest.level == "內部敏感"
    assert docx_text(res.files["docx"])[1] == "內部敏感：僅執行長與主管"
    assert res.files["md"].read_text(encoding="utf-8").startswith("> 內部敏感：僅執行長與主管")
    assert res.files["txt"].read_text(encoding="utf-8").startswith("【內部敏感：僅執行長與主管】")
    # 三個系統都上線時，結論是三句 headline
    assert res.digest.sentences[2] == "項目財務總覽：finance 的結論句。"


def test_unknown_level_treated_as_client_confidential(env):
    write(env["intel"], {**INTEL, "classification": "亂寫"})
    assert run(env["make"]()).digest.level == "客戶機密"


@pytest.mark.parametrize(
    "content, status_word",
    [
        (None, "本週未更新（找不到摘要檔）"),
        ("{壞掉", "摘要檔格式錯誤"),
        ({k: v for k, v in PROGRESS.items() if k != "headline"}, "摘要檔格式錯誤"),
        ({**PROGRESS, "system": "intel"}, "摘要檔格式錯誤"),
        ({**PROGRESS, "as_of": "2026-10-06"}, "本週未更新"),
        ({**PROGRESS, "as_of": "2026-10-20"}, "摘要檔日期晚於週報日期"),
    ],
)
def test_missing_or_bad_system_still_produces_digest(env, content, status_word):
    if content is None:
        env["progress"].unlink()
    else:
        write(env["progress"], content)
    res = run(env["make"]())
    d = res.digest
    assert status_word in d.sentences[0]
    assert all(
        m.this.startswith(status_word.split("（")[0])
        for m in d.metrics
        if m.key.startswith("progress.")
    )
    assert d.decisions == []  # 沒納入的系統，預警也不列
    txt = res.files["txt"].read_text(encoding="utf-8")
    assert status_word in txt and "外部情報摘要：intel 的結論句。" in txt
    assert any(status_word in x for x in d.data_quality)


def test_stale_boundary(env):
    write(env["progress"], {**PROGRESS, "as_of": "2026-10-07"})  # 剛好 7 天：仍算已更新
    assert run(env["make"]()).inputs[0].status == "ok"
    write(env["progress"], {**PROGRESS, "as_of": "2026-10-06"})  # 8 天
    assert run(env["make"]()).inputs[0].status == "stale"


def test_rerun_identical(env):
    a = run(env["make"]())
    h = {k: hashlib.sha256(p.read_bytes()).hexdigest() for k, p in a.files.items()}
    b = run(env["make"]())
    assert {k: hashlib.sha256(p.read_bytes()).hexdigest() for k, p in b.files.items()} == h


def test_truncates_to_fit_char_limit(env):
    long = summary(
        "progress",
        headline="很長" * 300,
        alerts=[
            {
                "level": "紅",
                "rule": "P01",
                "why": "原因很長" * 10,
                "subject": f"案{i}",
                "suggested_owner": "某人",
            }
            for i in range(30)
        ],
        data_quality=[f"問題{i}" * 10 for i in range(30)],
    )
    write(env["progress"], long)
    d = run(env["make"]()).digest
    assert d.body_size() <= 1400 or any("字上限" in x for x in d.missing)
    assert len(d.data_quality) < 5 and d.data_quality_more > 0


@pytest.mark.skipif(shutil.which("soffice") is None, reason="需要 LibreOffice 才能轉 PDF")
def test_docx_is_one_page_even_when_full(env, tmp_path):
    full = summary(
        "progress",
        headline="本週有 15 案需要注意：紅色 11 案（下一步日期已過 10 案、卡關 1 案），黃色 4 案；另有 1 位負責人案數偏多。",
        metrics=PROGRESS["metrics"],
        alerts=[
            {
                "level": "紅",
                "rule": "P01",
                "why": "下一步日期 2026-10-07 已過 7 天，階段仍是「進行中」",
                "subject": f"某客戶某專案名稱{i}",
                "suggested_owner": "某某（與某某）",
            }
            for i in range(12)
        ],
        data_quality=[
            "看板資料日期是 2026-10-05，距今 9 天，內容可能沒有更新，部分「已過日期」的案子可能其實已完成"
        ]
        * 8,
    )
    write(env["progress"], full)
    res = run(env["make"]())
    out = tmp_path / "pdf"
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out),
            str(res.files["docx"]),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    info = subprocess.run(
        ["pdfinfo", str(next(out.glob("*.pdf")))], capture_output=True, text=True
    ).stdout
    assert re.search(r"Pages:\s+1\b", info), info


def test_finance_path_from_env(env, monkeypatch, tmp_path):
    cfg = yaml.safe_load((env["cfg"] / "digest.yaml").read_text(encoding="utf-8"))
    cfg["systems"][2]["launched"] = True
    (env["cfg"] / "digest.yaml").write_text(
        yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8"
    )
    monkeypatch.delenv("FINANCE_OUT_DIR", raising=False)
    res = run(env["make"]())
    assert res.inputs[2].status == "missing" and "環境變數" in res.inputs[2].reason
    folder = tmp_path / "company"
    write(folder / "summary.json", summary("finance", level="內部敏感"))
    monkeypatch.setenv("FINANCE_OUT_DIR", str(folder))
    assert run(env["make"]()).inputs[2].status == "ok"


def test_model_links_only_from_intel_highlights(env):
    res = run(env["make"]())
    urls = set(re.findall(r"https?://\S+", res.files["txt"].read_text(encoding="utf-8")))
    assert urls == {f"https://example.com/{i}" for i in range(3)}


@pytest.mark.parametrize(
    "v, prev, out",
    [(0.3, 0.1, "+0.2"), (5, 5.0, "持平"), (None, 3, "待補"), (3, "x", "待補"), (True, 1, "待補")],
)
def test_change_format(v, prev, out):
    assert fmt_change(v, prev) == out


def test_digest_never_uses_llm():
    for p in (ROOT / "src" / "digest").glob("*.py"):
        assert "llm_client" not in p.read_text(encoding="utf-8"), p.name


def test_run_weekly_continues_after_failure(monkeypatch, capsys):
    import scripts.run_weekly as weekly

    calls = []

    def boom(argv):
        calls.append(("intel", argv))
        raise RuntimeError("x")

    def bad_args(argv):
        calls.append(("progress", argv))
        raise SystemExit(2)  # 例如參數錯誤

    def digest(argv):
        calls.append(("digest", argv))
        return 0

    monkeypatch.setattr(
        weekly,
        "STEPS",
        [("intel", "外部情報摘要", boom, ["--weekly"]), ("progress", "項目進度看板", bad_args, [])],
    )
    monkeypatch.setattr(weekly, "DIGEST", ("管理層週報", digest))
    assert weekly.main(["--as-of", "2026-10-14"]) == 1
    assert calls == [
        ("intel", ["--weekly", "--as-of", "2026-10-14"]),
        ("progress", ["--as-of", "2026-10-14"]),
        ("digest", ["--as-of", "2026-10-14", "--failed", "intel,progress"]),
    ]
    out = capsys.readouterr().out
    assert (
        "以下步驟未完成：外部情報摘要、項目進度看板（週報中已標示為「本週執行失敗，未納入」）"
        in out
    )


def test_failed_system_not_shown_with_old_numbers(env):
    res = run(env["make"](failed_systems={"progress"}))
    d = res.digest
    assert "本週執行失敗" in d.sentences[0]
    assert all(
        m.this.startswith("本週執行失敗") for m in d.metrics if m.key.startswith("progress.")
    )
    assert d.decisions == []


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), "12", True])
def test_non_finite_or_non_number_values_rejected(env, bad):
    data = json.loads(json.dumps(PROGRESS))
    data["metrics"][0]["value"] = bad
    env["progress"].write_text(
        json.dumps(data, ensure_ascii=False, allow_nan=True), encoding="utf-8"
    )
    res = run(env["make"]())  # 不崩潰
    assert res.inputs[0].status == "invalid"
    assert "格式錯誤" in res.digest.sentences[0]


def test_newlines_collapsed_and_long_text_cut(env):
    data = {
        **PROGRESS,
        "headline": "第一行\n## 假標題\n" * 40,
        "data_quality": ["多行\n" * 30] * 3,
        "alerts": [
            {
                "level": "紅",
                "rule": "P01",
                "why": "原因\n" * 50,
                "subject": "案",
                "suggested_owner": "人",
            }
        ],
    }
    write(env["progress"], data)
    d = run(env["make"]()).digest
    for t in d.sentences + d.decisions + d.data_quality:
        assert "\n" not in t
        assert len(t) < 400
    assert d.body_size() <= 1400


def test_md_has_no_injected_links(env):
    data = {
        **INTEL,
        "run_id": "[點我](javascript:alert(1))\n# 假標題",
        "highlights": [
            {"title": "正常", "url": "https://example.com/ok"},
            {"title": "壞", "url": "javascript:alert(1)"},
            {"title": "釣魚", "url": "https://x.com/a)[官方](https://evil.example"},
        ],
    }
    write(env["intel"], data)
    res = run(env["make"]())
    md = res.files["md"].read_text(encoding="utf-8")
    assert "](javascript" not in md and "\n# 假標題" not in md and "evil.example" not in md
    assert "<https://example.com/ok>" in md
    assert [i["title"] for i in res.digest.intel] == ["正常"]
    assert any(
        "連結不是 http" in x for x in res.digest.data_quality + [f"{res.digest.data_quality_more}"]
    )


@pytest.mark.parametrize(
    "mutate, msg",
    [
        (lambda c: c["systems"][2].update(launched="false"), "launched 只能寫 true 或 false"),
        (lambda c: c.pop("max_chars"), "max_chars 應為 1 以上的整數"),
        (lambda c: c.update(decisions_max=-1), "decisions_max 應為 1 以上的整數"),
        (lambda c: c["systems"][0].pop("name"), "缺少 key 或 name"),
        (lambda c: c["systems"].append(dict(c["systems"][0])), "重複"),
    ],
)
def test_bad_digest_config_is_plain_error(env, mutate, msg):
    from digest.sources import DigestConfigError

    p = env["cfg"] / "digest.yaml"
    c = yaml.safe_load(p.read_text(encoding="utf-8"))
    mutate(c)
    p.write_text(yaml.safe_dump(c, allow_unicode=True), encoding="utf-8")
    with pytest.raises(DigestConfigError, match=msg):
        run(env["make"]())


def test_bad_metrics_config_is_plain_error(env):
    from digest.sources import DigestConfigError

    p = env["cfg"] / "metrics.yaml"
    c = yaml.safe_load(p.read_text(encoding="utf-8"))
    del c["metrics"][0]["label"]
    p.write_text(yaml.safe_dump(c, allow_unicode=True), encoding="utf-8")
    with pytest.raises(DigestConfigError, match="缺少 key、label 或 system"):
        run(env["make"]())


def test_unknown_alert_level_and_classification_reported(env):
    write(
        env["progress"],
        {
            **PROGRESS,
            "classification": "機密吧",
            "alerts": [
                {"level": "紅色", "rule": "P01", "why": "x", "subject": "a"},
                {"level": "紅", "rule": "P01", "why": "y", "subject": "b"},
            ],
        },
    )
    d = run(env["make"]()).digest
    assert len(d.decisions) == 1
    assert any("燈號不是「紅」或「黃」" in x for x in d.data_quality)
    assert any("資料等級「機密吧」不認得" in x for x in d.missing)


def test_stale_days_null_is_noted(env):
    p = env["cfg"] / "thresholds.yaml"
    c = yaml.safe_load(p.read_text(encoding="utf-8"))
    c["digest"]["stale_days"]["value"] = None
    p.write_text(yaml.safe_dump(c, allow_unicode=True), encoding="utf-8")
    assert any("未設定" in x for x in run(env["make"]()).digest.missing)


def test_array_json_reason_and_run_id_has_input_hash(env):
    write(env["progress"], "[1, 2]")
    res = run(env["make"]())
    assert res.inputs[0].reason == "摘要檔內容不是一筆物件"
    rid = res.digest.run_id
    assert re.fullmatch(r"digest-20261014-[0-9a-f]{8}", rid)
    write(env["intel"], {**INTEL, "headline": "改了"})
    assert run(env["make"]()).digest.run_id != rid


@pytest.mark.parametrize("v, s", [(1e-7, "1E-7"), (0.5, "0.5"), (12, "12"), (1234567.0, "1234567")])
def test_value_format_no_scientific(v, s):
    from digest.build import fmt_value

    assert fmt_value(v) == s


@pytest.mark.parametrize("big", ["1" + "0" * 400, "-1" + "0" * 400, "1e400"])
def test_huge_numbers_do_not_crash(env, big):
    raw = json.dumps(PROGRESS, ensure_ascii=False).replace('"value": 11', f'"value": {big}', 1)
    env["progress"].write_text(raw, encoding="utf-8")
    res = run(env["make"]())
    assert res.inputs[0].status == "invalid" and set(res.files) == {"docx", "md", "txt"}


def test_urls_userinfo_rejected_parentheses_allowed(env):
    write(
        env["intel"],
        {
            **INTEL,
            "highlights": [
                {"title": "維基", "url": "https://zh.wikipedia.org/wiki/碳費_(臺灣)"},
                {"title": "假", "url": "https://good.com@evil.com/"},
            ],
        },
    )
    d = run(env["make"]()).digest
    assert [i["title"] for i in d.intel] == ["維基"]


def test_duplicate_metric_key_reported(env):
    write(env["progress"], {**PROGRESS, "metrics": PROGRESS["metrics"] + [PROGRESS["metrics"][0]]})
    d = run(env["make"]()).digest
    assert (
        any("progress.red_alerts 重複出現" in x for x in d.data_quality + d.missing)
        or d.data_quality_more
    )


def test_long_units_and_labels_cut(env):
    data = json.loads(json.dumps(PROGRESS))
    for m in data["metrics"]:
        m["unit"] = "單位" * 80
    write(env["progress"], data)
    d = run(env["make"]()).digest
    assert all(len(m.unit) <= 12 for m in d.metrics)


def test_run_id_changes_with_broken_file_and_log_has_inputs(env):
    write(env["progress"], "{壞1")
    a = run(env["make"]()).digest.run_id
    write(env["progress"], "{壞2")
    b = run(env["make"]()).digest.run_id
    assert a != b
    log = (env["tmp"] / "logs" / f"run_{b}.log").read_text(encoding="utf-8")
    assert "程式版本" in log and "progress.json" in log


@pytest.mark.skipif(shutil.which("soffice") is None, reason="需要 LibreOffice 才能轉 PDF")
def test_three_systems_extreme_still_one_page(env, tmp_path):
    cfg = env["cfg"]
    d = yaml.safe_load((cfg / "digest.yaml").read_text(encoding="utf-8"))
    d["systems"][2]["launched"] = True
    (cfg / "digest.yaml").write_text(yaml.safe_dump(d, allow_unicode=True), encoding="utf-8")
    ms = [
        {"key": f"{s}.m{i}", "label": "很長的指標名稱" * 10, "system": s}
        for s in ("progress", "intel", "finance")
        for i in range(3)
    ][:8]
    m = yaml.safe_load((cfg / "metrics.yaml").read_text(encoding="utf-8"))
    m["metrics"] = ms
    (cfg / "metrics.yaml").write_text(yaml.safe_dump(m, allow_unicode=True), encoding="utf-8")

    def full(key):
        return summary(
            key,
            level="內部敏感",
            headline="非常長的結論" * 40,
            metrics=[
                {
                    "key": f"{key}.m{i}",
                    "label": "x",
                    "value": 123456789,
                    "unit": "單位" * 80,
                    "prev_value": 1,
                }
                for i in range(3)
            ],
            alerts=[
                {
                    "level": "紅",
                    "rule": "R",
                    "why": "原因" * 80,
                    "subject": "案" * 50,
                    "suggested_owner": "人" * 30,
                }
            ]
            * 50,
            data_quality=["品質" * 80] * 50,
            highlights=[{"title": "標題" * 200, "url": "https://example.com/x"}] * 10,
        )

    paths = {
        k: write(tmp_path / f"{k}_full.json", full(k)) for k in ("progress", "intel", "finance")
    }
    res = run(env["make"](overrides=paths))
    assert res.digest.body_size() <= 1400
    out = tmp_path / "pdf"
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out),
            str(res.files["docx"]),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    info = subprocess.run(
        ["pdfinfo", str(next(out.glob("*.pdf")))], capture_output=True, text=True
    ).stdout
    assert re.search(r"Pages:\s+1\b", info), info
