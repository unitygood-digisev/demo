"""階段 2b：會議紀錄 → 看板補齊建議（全部用假資料）。"""

import datetime as dt
import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from progress.backfill import BackfillConfig, run
from progress.board import ProgressDataError
from progress.meetings import extract_dates, match_cases, parse_meeting

from .conftest import ROOT, SAMPLE, sample_board

MEET = SAMPLE / "meetings"
ALIASES = SAMPLE / "meeting_aliases_sample.yaml"
AS_OF = dt.date(2026, 10, 15)
BOARD = sample_board(dt.date(2026, 10, 14))


@pytest.fixture
def meet_dir(tmp_path):
    d = tmp_path / "meetings"
    shutil.copytree(MEET, d)
    return d


@pytest.fixture
def make(tmp_path, config_dir, meet_dir):
    def _make(**kw):
        base = dict(
            as_of=AS_OF,
            board_path=BOARD,
            meetings_dir=meet_dir,
            aliases_path=ALIASES,
            out_dir=tmp_path / "out",
            config_dir=config_dir,
            log_dir=tmp_path / "logs",
        )
        base.update(kw)
        return BackfillConfig(**base)

    return _make


def sheet(path: Path, name: str) -> list[dict]:
    ws = load_workbook(path)[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[1])  # 第 1 列是資料等級，第 2 列是表頭
    out = []
    for r in rows[2:]:
        if all(c is None for c in r):  # 表格之後是頁尾
            break
        out.append(dict(zip(hdr, r, strict=False)))
    return out


def all_text(path: Path) -> str:
    wb = load_workbook(path)
    return "\n".join(
        str(c) for ws in wb for r in ws.iter_rows(values_only=True) for c in r if c is not None
    )


# ---- 日期 ----
@pytest.mark.parametrize(
    "text, exact, approx",
    [
        ("預計 10/20 訪談", ["2026-10-20"], []),
        ("10 月底完成", [], ["2026年10月底"]),
        ("2026年11月3日開會", ["2026-11-03"], []),
        ("11 月 5 日提案", ["2026-11-05"], []),
        ("展期：11/11～11/13", ["2026-11-11", "2026-11-13"], []),
        ("暫緩到明年 3 月", [], ["2027年3月"]),
        ("年底前完成", [], ["2026年底"]),
        ("達成率 80%、S1、30×30、7300家", [], []),
        ("13/45 不是日期", [], []),
        ("12 月中旬", [], ["2026年12月中旬"]),
    ],
)
def test_extract_dates(text, exact, approx):
    e, a = extract_dates(text, dt.date(2026, 10, 1))
    assert [d.isoformat() for d in e] == exact
    assert [x.label for x in a] == approx


# ---- 對應案名 ----
CASES = [
    "（假資料）甲案",
    "甲機關 標章交流會",
    "甲機關 遊程工作坊第二次",
    "乙院 SROI",
    "專班：乙院醫院",
]


@pytest.mark.parametrize(
    "text, expected",
    [
        ("甲案：交稿", {"（假資料）甲案"}),
        ("甲機關專案：8 月底", {"甲機關 標章交流會", "甲機關 遊程工作坊第二次"}),  # 不猜
        ("乙院醫院 SROI：9/10", {"乙院 SROI", "專班：乙院醫院"}),
        ("乙院醫院 SROI：9/10 結案", {"乙院 SROI"}),  # 有對應表時以對應表為準
        ("醫院工作坊", set()),  # 籠統字詞不拿來比對
    ],
)
def test_match_cases(text, expected):
    aliases = {"乙院醫院 SROI": "乙院 SROI"} if "結案" in text else {}
    generic = ["假資料", "醫院", "工作坊", "SROI", "專班"]
    assert match_cases(text, CASES, aliases, generic) == expected


# ---- 拆解會議紀錄 ----
def test_parse_skips_finance_hr_admin(config_dir):
    from progress.meetings import MeetingRules

    m = parse_meeting(MEET / "業務會議紀錄_假資料_20261001.txt", MeetingRules.load(config_dir))
    assert m.date == dt.date(2026, 10, 1)
    skipped = {ln.no for ln in m.lines if ln.skipped}
    text = (MEET / "業務會議紀錄_假資料_20261001.txt").read_text(encoding="utf-8").splitlines()
    for i, t in enumerate(text, start=1):
        if any(w in t for w in ("營收", "人力", "子案")):
            assert i in skipped


def test_filename_without_date_is_loud(make, meet_dir):
    (meet_dir / "沒有日期.txt").write_text("* 己案： 10/20", encoding="utf-8")
    res = run(make())
    files = sheet(res.path, "會議紀錄清單")
    assert any("檔名沒有日期" in str(r["狀態"]) for r in files)
    assert "沒有日期.txt" in res.headline or "無法使用" in res.headline


def test_non_utf8_is_listed_not_crash(make, meet_dir):
    (meet_dir / "業務會議紀錄_20261002.txt").write_bytes("* 己案： 10/20".encode("big5"))
    res = run(make())
    files = sheet(res.path, "會議紀錄清單")
    assert any("UTF-8" in str(r["狀態"]) for r in files)


def test_date_in_text_differs_from_filename(make, meet_dir):
    (meet_dir / "業務會議紀錄_20261003.txt").write_text(
        "* 會議日期及時間：2026年9月3日\n* 己案： 10/20\n", encoding="utf-8"
    )
    res = run(make())
    files = {r["檔案"]: r["狀態"] for r in sheet(res.path, "會議紀錄清單")}
    assert "和檔名不符" in files["業務會議紀錄_20261003.txt"]


def test_no_meeting_files_stops(make, tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(ProgressDataError, match="沒有會議紀錄"):
        run(make(meetings_dir=empty))


# ---- 端到端 ----
def test_end_to_end_sample(make):
    res = run(make())
    fill = {r["案子"]: r for r in sheet(res.path, "建議補上的日期")}
    assert set(fill) == {"（假資料）己案", "（假資料）丑案"}
    assert fill["（假資料）己案"]["建議日期"] == "2026-10-20"
    assert "2026年10月底" in fill["（假資料）己案"]["約略時程"]
    assert fill["（假資料）丑案"]["建議日期"] == "待補"  # 只有約略時程，不換算成某一天
    assert "2027年3月" in fill["（假資料）丑案"]["約略時程"]
    assert "暫緩" in fill["（假資料）丑案"]["紀錄提到的階段字眼"]
    # 每條建議都追得到出處
    assert fill["（假資料）己案"]["出處"] == "業務會議紀錄_假資料_20261001.txt 第 8 行"

    diff = {r["案子"]: r for r in sheet(res.path, "與看板不一致")}
    assert set(diff) == {"（假資料）壬案", "（假資料）辰案", "（假資料）乙案"}
    assert diff["（假資料）壬案"]["看板"] == "2026-11-05"
    assert "2026-11-12" in diff["（假資料）壬案"]["會議紀錄"]
    assert "2026-11-20" in diff["（假資料）辰案"]["會議紀錄"]  # 標題「辰案相關」下的條目
    assert "已結案" in diff["（假資料）乙案"]["會議紀錄"]  # 透過對應表

    missing = {r["紀錄中的名稱"]: r for r in sheet(res.path, "看板沒有的案子")}
    assert set(missing) == {"新客戶A"}
    assert "2026-11-05" in missing["新客戶A"]["提到的日期"]

    manual = sheet(res.path, "待人工確認")
    reasons = " ".join(str(r["原因"]) for r in manual)
    assert "同名" in reasons  # 甲案在看板上重複
    assert "對到多個案子" in reasons  # 寅案與卯案
    assert "對不到案子" in reasons  # 11/25 內部分享會
    # 已完成的案子、日期一致的案子、只存檔的舊紀錄都不出現
    text = all_text(res.path)
    assert "（假資料）庚案" not in text
    assert "（假資料）卯案" not in {r["案子"] for r in sheet(res.path, "與看板不一致")}
    assert "舊客戶B" not in text
    files = {r["檔案"]: r["狀態"] for r in sheet(res.path, "會議紀錄清單")}
    assert "只存檔" in files["業務會議紀錄_假資料_20260801.txt"]
    assert res.headline.startswith("會議紀錄可以幫看板補上 2 案的日期")


def test_no_raw_text_or_finance_in_output(make):
    res = run(make())
    text = all_text(res.path)
    for word in (
        "營收",
        "達成率",
        "未開發票",
        "300 萬",
        "招募",
        "面試",
        "發票",
        "訪談",
        "排課",
        "感謝函",
        "聯合論壇",
        "內部分享會",
        "收資料",
    ):
        assert word not in text, word
    wb = load_workbook(res.path)
    assert "內部一般" in str(wb["總覽"]["A1"].value)


def test_inputs_untouched_and_rerun_identical(make, meet_dir):
    def digest():
        h = hashlib.sha256(BOARD.read_bytes())
        for p in sorted(meet_dir.iterdir()):
            h.update(p.read_bytes())
        return h.hexdigest()

    before = digest()
    a = run(make()).path.read_bytes()
    b = run(make()).path.read_bytes()
    assert a == b
    assert digest() == before


def test_threshold_from_config(make, config_dir):
    import yaml

    p = config_dir / "thresholds.yaml"
    d = yaml.safe_load(p.read_text(encoding="utf-8"))
    d["meetings"]["max_age_days"]["value"] = 100  # 8/1 的紀錄也納入
    p.write_text(yaml.safe_dump(d, allow_unicode=True), encoding="utf-8")
    res = run(make())
    assert "舊客戶B" in {r["紀錄中的名稱"] for r in sheet(res.path, "看板沒有的案子")}


@pytest.mark.parametrize("bad", [None, -1, "45", True])
def test_bad_threshold_stops(make, config_dir, bad):
    import yaml

    p = config_dir / "thresholds.yaml"
    d = yaml.safe_load(p.read_text(encoding="utf-8"))
    d["meetings"]["max_age_days"]["value"] = bad
    p.write_text(yaml.safe_dump(d, allow_unicode=True), encoding="utf-8")
    with pytest.raises(ProgressDataError, match="max_age_days"):
        run(make())


def test_missing_alias_file_is_noted(make, tmp_path):
    res = run(make(aliases_path=tmp_path / "none.yaml"))
    assert any("對應表" in n for n in res.notes)


def test_alias_to_unknown_case_is_loud(make, tmp_path):
    p = tmp_path / "a.yaml"
    p.write_text("aliases:\n  簡稱乙: 看板上沒有的案名\n", encoding="utf-8")
    res = run(make(aliases_path=p))
    reasons = " ".join(str(r["原因"]) for r in sheet(res.path, "待人工確認"))
    assert "對應表" in reasons


def test_meeting_after_board_is_used_and_marked(make, meet_dir):
    (meet_dir / "業務會議紀錄_20261020.txt").write_text("* 己案： 10/28 訪談\n", encoding="utf-8")
    res = run(make())
    files = {r["檔案"]: r["狀態"] for r in sheet(res.path, "會議紀錄清單")}
    assert "比看板新" in files["業務會議紀錄_20261020.txt"]
    fill = {r["案子"]: r for r in sheet(res.path, "建議補上的日期")}
    assert "2026-10-28" in fill["（假資料）己案"]["其他提到的日期"]


def test_does_not_import_llm():
    src = (ROOT / "src" / "progress" / "meetings.py").read_text(encoding="utf-8")
    src += (ROOT / "src" / "progress" / "backfill.py").read_text(encoding="utf-8")
    assert "llm_client" not in src


def test_real_aliases_file_is_gitignored():
    import subprocess

    r = subprocess.run(
        ["git", "check-ignore", "config/meeting_aliases.yaml", "data/raw/meetings/x.txt"],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    assert r.stdout.split() == ["config/meeting_aliases.yaml", "data/raw/meetings/x.txt"]


def test_label_only_line_is_context(make, meet_dir):
    (meet_dir / "業務會議紀錄_20261005.txt").write_text(
        "* 簡稱乙：\n   * 日期： 10/25 驗收\n   * 另外 10/27 回報\n", encoding="utf-8"
    )
    res = run(make())
    names = {r["紀錄中的名稱"] for r in sheet(res.path, "看板沒有的案子")}
    assert "日期" not in names  # 「日期：」屬於上一行的案子，不是新案子
    diff = {r["案子"]: r for r in sheet(res.path, "與看板不一致")}
    assert "2026-10-25" in diff["（假資料）乙案"]["會議紀錄"]


def test_short_name_inside_board_name_and_tai():
    cases = ["丙慈善醫院", "台北丁中心"]
    assert match_cases("丙慈善", cases, {}, ["醫院"]) == {"丙慈善醫院"}
    assert match_cases("臺北丁中心：9/21", cases, {}, ["醫院"]) == {"台北丁中心"}


def test_ignore_list_and_sentence_labels(make, meet_dir, tmp_path):
    (meet_dir / "業務會議紀錄_20261006.txt").write_text(
        "* 現場小卡： 共 170 張\n* 已結案，請款中： 8/24 開立\n", encoding="utf-8"
    )
    p = tmp_path / "a.yaml"
    p.write_text("aliases:\n  簡稱乙: （假資料）乙案\nignore:\n  - 現場小卡\n", encoding="utf-8")
    res = run(make(aliases_path=p))
    names = {r["紀錄中的名稱"] for r in sheet(res.path, "看板沒有的案子")}
    assert "現場小卡" not in names and not any("請款" in n for n in names)


def test_vague_done_words_not_flagged(make, meet_dir):
    # 「預計 11 月完成」「已完成第一版」不是結案
    (meet_dir / "業務會議紀錄_20261007.txt").write_text(
        "* 壬案： 已完成第一版，預計 11 月完成\n", encoding="utf-8"
    )
    res = run(make())
    diff = {r["案子"]: r for r in sheet(res.path, "與看板不一致")}
    assert "結案" not in diff.get("（假資料）壬案", {}).get("不一致", "")
