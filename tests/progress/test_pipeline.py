"""系統 1 端到端測試（用 data/sample/progress 的假看板）。"""

import datetime as dt
import hashlib
import json

import pytest
from openpyxl import Workbook, load_workbook

from progress.board import ProgressDataError
from progress.pipeline import run

from .conftest import AS_OF, SAMPLE, WEEKS, sample_board, set_threshold


def sheet_rows(path, name):
    ws = load_workbook(path)[name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1]
    out = []
    for r in rows[2:]:
        if not r or all(c is None for c in r):
            break
        out.append(dict(zip(header, r, strict=False)))
    return out


def all_cells(path):
    return " ".join(
        str(c)
        for ws in load_workbook(path)
        for r in ws.iter_rows(values_only=True)
        for c in r
        if c is not None
    )


def test_three_weeks_end_to_end(three_weeks):
    res = three_weeks()
    d = res.data
    assert d["metrics"] == {
        "red_projects": 3,
        "yellow_projects": 3,
        "green_projects": 15,
        "active_projects": 17,
        "missing_next_date": 1,
        "total_projects": 21,
    }
    got = sorted((a["rule"], a["subject"]) for a in d["alerts"])
    assert got == [
        ("P01", "（假資料）甲案"),
        ("P01", "（假資料）辛案"),
        ("P02", "（假資料）丙案"),
        ("P02", "（假資料）乙案"),
        ("P03", "（假資料）戊案"),
        ("P04", "（假資料）壬案"),
        ("P05", "顧問甲"),
    ]
    ch = d["changes"]
    assert ch["prev_date"] == "2026-10-07"
    assert {(a["rule"], a["subject"]) for a in ch["resolved"]} == {
        ("P01", "（假資料）庚案"),
        ("P02", "（假資料）辛案"),
    }
    assert ch["removed_projects"] == ["（假資料）亥案"]
    md = res.files["md"].read_text(encoding="utf-8")
    assert md.splitlines()[0] == "> 資料等級：內部一般"
    assert (
        "本週有 6 案需要注意：紅色 3 案（下一步日期已過 2 案、卡關 1 案），黃色 3 案；另有 1 位負責人案數偏多。"
        in md
    )
    assert "新增 4 項預警、解除 2 項、持續 3 項" in md


def test_overview_numbers_match_alert_list(three_weeks):
    res = three_weeks()
    alerts = sheet_rows(res.files["xlsx"], "預警清單")
    projects = sheet_rows(res.files["xlsx"], "項目明細")
    ws = load_workbook(res.files["xlsx"])["總覽"]
    ov = {
        r[1]: r[2] for r in ws.iter_rows(values_only=True) if r and r[1] and isinstance(r[2], int)
    }
    for light, label in (
        ("紅", "需要立即確認的案子"),
        ("黃", "需要留意的案子"),
        ("綠", "沒有預警的案子"),
    ):
        assert ov[label] == sum(1 for p in projects if p["燈號"] == light)
    for rule in ("P01", "P02", "P03", "P04", "P05"):
        n = next(v for k, v in ov.items() if k.startswith(rule))
        assert n == sum(1 for a in alerts if a["規則"] == rule)
    # 每個紅／黃案子都能在預警清單找到原因
    alerted = {a["案子／負責人"] for a in alerts}
    assert all(p["案子"] in alerted for p in projects if p["燈號"] != "綠")


def test_excluded_columns_never_in_output(three_weeks):
    res = three_weeks()
    text = (
        all_cells(res.files["xlsx"])
        + res.files["md"].read_text(encoding="utf-8")
        + res.files["summary"].read_text(encoding="utf-8")
    )
    assert "不續約" not in text and "備註內容" not in text


def test_data_quality_listed_not_crash(three_weeks):
    res = three_weeks()
    q = sheet_rows(res.files["xlsx"], "資料品質")
    got = {(r["來源列號"], r["欄位"], r["問題"].split("（")[0]) for r in q}
    assert (5, "客戶／案子", "案名重複，無法和上週比較，請改成不同名稱") in got
    assert (23, "客戶／案子", "案名重複，無法和上週比較，請改成不同名稱") in got
    assert (22, "下一個日期", "日期錯誤") in got
    assert (24, "階段", "不在允許值內") in got
    assert (25, "客戶／案子", "必填欄位空白") in got
    assert (26, "負責人", "進行中的案子沒有負責人") in got
    assert (10, "下一個日期", "進行中的案子沒有下一個日期，無法判斷是否延誤") in got
    # 日期寫錯的列不重複列「沒有下一個日期」
    assert sum(1 for r in q if r["來源列號"] == 22) == 1


def test_rerun_identical_bytes(make_cfg):
    a = run(make_cfg())
    h1 = {k: hashlib.sha256(v.read_bytes()).hexdigest() for k, v in a.files.items()}
    b = run(make_cfg())
    assert {k: hashlib.sha256(v.read_bytes()).hexdigest() for k, v in b.files.items()} == h1


def test_first_run_has_no_comparison(make_cfg):
    res = run(make_cfg())
    assert res.data["changes"] == {"available": False}
    assert "這是第一次執行" in res.files["md"].read_text(encoding="utf-8")
    assert any("快照不足" in n for n in res.data["notes"])


def test_threshold_change_only_via_config(make_cfg, config_dir):
    set_threshold(config_dir, "P02_due_soon_days", 0)  # 只剩今天到期
    res = run(make_cfg())
    assert sorted(a["subject"] for a in res.data["alerts"] if a["rule"] == "P02") == [
        "（假資料）乙案"
    ]
    set_threshold(config_dir, "P05_max_active_projects", 10)
    res = run(make_cfg())
    assert not [a for a in res.data["alerts"] if a["rule"] == "P05"]


def test_backfill_refused(make_cfg):
    run(make_cfg())
    with pytest.raises(ProgressDataError, match="早於已處理過的最新日期"):
        run(make_cfg(as_of=WEEKS[1], input_path=sample_board(WEEKS[1])))


def test_board_date_and_stale_note(make_cfg):
    assert (
        run(make_cfg(as_of=AS_OF + dt.timedelta(days=7), input_path=sample_board(AS_OF))).data[
            "board_note"
        ]
        is None
    )
    res = run(make_cfg(as_of=AS_OF + dt.timedelta(days=8), input_path=sample_board(AS_OF)))
    assert res.data["board_date"] == "2026-10-14"
    assert "距今 8 天" in res.data["board_note"]
    assert "注意：看板資料日期是 2026-10-14" in res.files["md"].read_text(encoding="utf-8")
    assert run(make_cfg(as_of=AS_OF + dt.timedelta(days=9), input_path=sample_board(AS_OF))).data[
        "board_note"
    ]


def test_summary_json_for_digest(three_weeks):
    res = three_weeks()
    s = json.loads(res.files["summary"].read_text(encoding="utf-8"))
    assert s["system"] == "progress" and s["classification"] == "內部一般"
    m = {x["key"]: x for x in s["metrics"]}
    assert m["progress.red_alerts"]["value"] == 3 and m["progress.red_alerts"]["prev_value"] == 2
    assert {a["level"] for a in s["alerts"]} == {"紅", "黃"}


def _board(path, header_row=4, header=None, rows=(), sheet="項目看板"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for _ in range(header_row - 1):
        ws.append(["說明"])
    ws.append(
        header
        or [
            "客戶／案子",
            "類別",
            "階段",
            "負責人",
            "近期下一步",
            "下一個日期",
            "卡點／風險",
            "備註",
        ]
    )
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def test_header_found_anywhere_and_row_numbers_real(make_cfg, tmp_path):
    p = _board(
        tmp_path / "b.xlsx", header_row=7, rows=[["A", "x", "進行中", "甲", "s", "2026/13/1"]]
    )
    res = run(make_cfg(input_path=p))
    q = sheet_rows(res.files["xlsx"], "資料品質")
    assert any(r["來源列號"] == 8 and r["欄位"] == "下一個日期" for r in q)


@pytest.mark.parametrize(
    "kind, expect",
    [
        ("no_sheet", "找不到分頁"),
        ("no_header", "找不到標題列"),
        ("empty", "檔案沒有資料列"),
        ("missing_date_col", "缺少欄位（必填）：下一個日期"),
        ("missing_owner_col", "缺少欄位（必填）：負責人"),
    ],
)
def test_unreadable_board_stops_without_writing(make_cfg, tmp_path, kind, expect):
    run(make_cfg(as_of=WEEKS[0], input_path=sample_board(WEEKS[0])))
    before = sorted(p.name for p in (tmp_path / "interim").rglob("*"))
    p = tmp_path / "b.xlsx"
    header = ["客戶／案子", "類別", "階段", "負責人", "近期下一步", "下一個日期"]
    if kind == "no_sheet":
        _board(p, sheet="別的分頁", rows=[["A", "x", "進行中", "甲", "s", None]])
    elif kind == "no_header":
        _board(p, header=["案名", "階段"], rows=[["A", "進行中"]])
    elif kind == "empty":
        _board(p)
    elif kind == "missing_date_col":
        _board(
            p,
            header=[h if h != "下一個日期" else "下一個 日期" for h in header],
            rows=[["A", "x", "進行中", "甲", "s", None]],
        )
    else:
        _board(
            p, header=[h for h in header if h != "負責人"], rows=[["A", "x", "進行中", "s", None]]
        )
    with pytest.raises(ProgressDataError, match=expect):
        run(make_cfg(as_of=WEEKS[1], input_path=p))
    assert sorted(p.name for p in (tmp_path / "interim").rglob("*")) == before
    assert not list((tmp_path / "out").glob("*20261007*"))


def test_non_xlsx_input_refused(make_cfg, tmp_path):
    p = tmp_path / "項目進度看板.csv"
    p.write_text("說明\n客戶／案子,階段\nA,進行中\n", encoding="utf-8")
    with pytest.raises(ProgressDataError, match="必須是 .xlsx"):
        run(make_cfg(input_path=p))


def test_header_picks_real_header_row(make_cfg, tmp_path):
    p = tmp_path / "b.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "項目看板"
    ws.append(["客戶／案子"])  # 說明文字剛好也寫了錨點
    ws.append(["來源：2026/10/14 會議紀錄"])
    ws.append(["客戶／案子", "類別", "階段", "負責人", "近期下一步", "下一個日期"])
    ws.append(["A", "x", "卡關", "甲", "s", None])
    wb.save(p)
    res = run(make_cfg(input_path=p))
    assert res.data["metrics"]["total_projects"] == 1 and res.data["alerts"][0]["rule"] == "P03"
    assert res.data["board_date"] == "2026-10-14"


def test_missing_input_is_loud(make_cfg, tmp_path):
    with pytest.raises(ProgressDataError, match="找不到"):
        run(make_cfg(input_path=None))
    with pytest.raises(ProgressDataError, match="找不到看板檔案"):
        run(make_cfg(input_path=tmp_path / "nope.xlsx"))


def test_find_latest_input_in_raw(make_cfg, tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    for d in WEEKS:
        (raw / f"項目進度看板_{d.strftime('%Y%m%d')}.xlsx").write_bytes(
            sample_board(d).read_bytes()
        )
    res = run(make_cfg(input_path=None))
    assert res.data["input_file"] == "項目進度看板_20261014.xlsx"


def test_raw_board_not_modified(make_cfg):
    p = sample_board(AS_OF)
    before = hashlib.sha256(p.read_bytes()).hexdigest()
    run(make_cfg())
    assert hashlib.sha256(p.read_bytes()).hexdigest() == before


def test_corrupt_history_is_loud(make_cfg, tmp_path):
    run(make_cfg(as_of=WEEKS[0], input_path=sample_board(WEEKS[0])))
    (tmp_path / "interim" / "runs" / "20260930.json").write_text("{壞", encoding="utf-8")
    with pytest.raises(ProgressDataError, match="20260930.json"):
        run(make_cfg())


def test_sample_is_marked_fake():
    from openpyxl import load_workbook as lw

    for d in WEEKS:
        ws = lw(sample_board(d))["項目看板"]
        names = [r[0] for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
        assert names and all(n.startswith("（假資料）") for n in names)
    assert "假資料" in (SAMPLE.parent / "README.md").read_text(encoding="utf-8")


@pytest.mark.parametrize("value", ["七", "7", -1, True, 1.5])
def test_bad_threshold_is_plain_error(make_cfg, config_dir, value):
    set_threshold(config_dir, "P02_due_soon_days", value)
    with pytest.raises(ProgressDataError, match="progress.P02_due_soon_days 應為 0 以上的整數"):
        run(make_cfg())


@pytest.mark.parametrize(
    "content",
    ['{"metrics": {}, "projects": []}', "[1, 2]", '{"alerts": [1], "metrics": {}, "projects": []}'],
)
def test_runs_with_wrong_structure_is_loud(make_cfg, tmp_path, content):
    run(make_cfg(as_of=WEEKS[0], input_path=sample_board(WEEKS[0])))
    (tmp_path / "interim" / "runs" / "20260930.json").write_text(content, encoding="utf-8")
    with pytest.raises(ProgressDataError, match="20260930.json"):
        run(make_cfg())


def test_other_formats_in_raw_are_reported(make_cfg, tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "項目進度看板_20261007.xlsx").write_bytes(sample_board(WEEKS[1]).read_bytes())
    (raw / "項目進度看板_20261014.xls").write_bytes(b"old excel")
    (raw / "~$項目進度看板_20261014.xlsx").write_bytes(b"lock")
    res = run(make_cfg(input_path=None))
    assert res.data["input_file"] == "項目進度看板_20261007.xlsx"
    assert any("項目進度看板_20261014.xls" in n and "已略過" in n for n in res.data["notes"])


def test_formula_like_text_stays_text(make_cfg, tmp_path):
    p = _board(
        tmp_path / "b.xlsx", rows=[['=HYPERLINK("http://x")', "x", "卡關", "甲", "=1+1", None]]
    )
    # 模擬使用者在看板輸入以「=」開頭的文字（儲存格是文字型態，不是公式）
    wb0 = load_workbook(p)
    for c in wb0["項目看板"][5]:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"
    wb0.save(p)
    res = run(make_cfg(input_path=p))
    wb = load_workbook(res.files["xlsx"])
    cells = [
        c
        for ws in wb
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert cells and all(c.data_type == "s" for c in cells)


def test_duplicate_alerts_counted_in_changes(make_cfg, tmp_path):
    rows = [["A", "x", "卡關", "甲", "s", None], ["B", "x", "進行中", "甲", "s", None]]
    run(make_cfg(as_of=WEEKS[0], input_path=_board(tmp_path / "w1.xlsx", rows=rows)))
    rows2 = [["A", "x", "卡關", "甲", "s", None], ["A", "x", "卡關", "乙", "t", None]]
    res = run(make_cfg(as_of=WEEKS[1], input_path=_board(tmp_path / "w2.xlsx", rows=rows2)))
    assert len([a for a in res.data["alerts"] if a["rule"] == "P03"]) == 2
    ch = res.data["changes"]
    assert len(ch["new"]) == 1 and len(ch["ongoing"]) == 1


def test_same_day_different_board_is_noted(make_cfg, tmp_path):
    run(make_cfg())
    other = tmp_path / "另一份.xlsx"
    other.write_bytes(sample_board(AS_OF).read_bytes())
    res = run(make_cfg(input_path=other))
    assert any("今天稍早已用另一份看板" in n for n in res.data["notes"])


def test_board_date_prefers_source_line_and_future_date(make_cfg, tmp_path):
    p = tmp_path / "b.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "項目看板"
    ws.append(["規則自 2025/1/1 起適用"])
    for _ in range(14):
        ws.append(["說明"])
    ws.append(["來源：2026/10/20 會議紀錄"])  # 第 16 列，超過舊版只看 10 列的範圍
    ws.append(["客戶／案子", "類別", "階段", "負責人", "近期下一步", "下一個日期"])
    ws.append(["A", "x", "進行中", "甲", "s", None])
    wb.save(p)
    res = run(make_cfg(input_path=p))
    assert res.data["board_date"] == "2026-10-20"
    assert any("晚於執行日期" in n for n in res.data["notes"])


def test_null_thresholds_are_noted(make_cfg, config_dir):
    for k in ("P02_due_soon_days", "P05_max_active_projects"):
        set_threshold(config_dir, k, None)
    notes = run(make_cfg()).data["notes"]
    assert any("P02 門檻未設定" in n for n in notes) and any("P05 門檻未設定" in n for n in notes)


def test_missing_threshold_key_is_plain_error(make_cfg, config_dir):
    import yaml

    p = config_dir / "thresholds.yaml"
    data = yaml.safe_load(p.read_text(encoding="utf-8"))
    del data["progress"]["P05_max_active_projects"]
    p.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    with pytest.raises(ProgressDataError, match="缺少 progress.P05_max_active_projects"):
        run(make_cfg())


def test_uppercase_extension_in_raw(make_cfg, tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "項目進度看板_20261014.XLSX").write_bytes(sample_board(AS_OF).read_bytes())
    assert run(make_cfg(input_path=None)).data["input_file"] == "項目進度看板_20261014.XLSX"


def test_board_date_in_roc_year(make_cfg, tmp_path):
    p = tmp_path / "b.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "項目看板"
    ws.append(["來源：115/10/12 會議紀錄"])
    ws.append(["客戶／案子", "類別", "階段", "負責人", "近期下一步", "下一個日期"])
    ws.append(["A", "x", "進行中", "甲", "s", None])
    wb.save(p)
    assert run(make_cfg(input_path=p)).data["board_date"] == "2026-10-12"


def test_pipe_in_name_has_no_stray_backslash(make_cfg, tmp_path):
    p = _board(tmp_path / "b.xlsx", rows=[["A|B 案", "x", "卡關", "甲", "s", None]])
    md = run(make_cfg(input_path=p)).files["md"].read_text(encoding="utf-8")
    assert "A／B 案" in md and "\\／" not in md


def test_refusal_leaves_no_log(make_cfg, tmp_path):
    run(make_cfg(as_of=WEEKS[0], input_path=sample_board(WEEKS[0])))
    (tmp_path / "interim" / "runs" / "20260930.json").write_text("[1]", encoding="utf-8")
    logs_before = sorted(p.name for p in (tmp_path / "logs").glob("*"))
    with pytest.raises(ProgressDataError):
        run(make_cfg())
    assert sorted(p.name for p in (tmp_path / "logs").glob("*")) == logs_before
