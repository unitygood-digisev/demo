"""獨立驗收找出的情況：回補日期、後續報導、markdown、追蹤表改動、中間檔損壞。"""

import datetime as dt
import json

import pytest
from openpyxl import load_workbook

from intel.output import md_escape
from intel.pipeline import run
from intel.store import IntelDataError

from .conftest import AS_OF, SAMPLE
from .test_pipeline import all_output_text, tracking_rows

FEED = '<?xml version="1.0"?><rss version="2.0"><channel><title>t</title>{items}</channel></rss>'
ITEM = "<item><title>{t}</title><link>{u}</link><pubDate>{d}</pubDate><description>{x}</description></item>"


def write_feed(folder, name, items):
    p = folder / f"{name}.xml"
    p.write_text(FEED.format(items="".join(ITEM.format(**i) for i in items)), encoding="utf-8")
    return p


def write_sources(folder, names):
    p = folder / "s.yaml"
    lines = ["sources:"]
    for n in names:
        lines.append(
            f"  - {{id: {n}, 名稱: 來源{n}, 類型: rss, 網址: '{folder / (n + '.xml')}', 分類: 其他}}"
        )
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return p


def test_backfill_refused_and_nothing_lost(make_cfg):
    a = run(make_cfg())
    md = a.files["daily"].read_text(encoding="utf-8")
    with pytest.raises(IntelDataError, match="早於已處理過的最新日期"):
        run(make_cfg(as_of=AS_OF - dt.timedelta(days=1)))
    b = run(make_cfg())
    assert b.files["daily"].read_text(encoding="utf-8") == md
    assert len(tracking_rows(b.files["tracking"])) == 8


def test_new_report_of_old_event_is_listed(make_cfg, tmp_path):
    feeds = tmp_path / "feeds"
    feeds.mkdir()
    old = {
        "t": "Board issues amendments to climate disclosure standard",
        "u": "https://a.example.com/1",
        "d": "Mon, 12 Oct 2026 09:00:00 +0000",
        "x": "climate disclosure news",
    }
    write_feed(feeds, "a", [old])
    write_feed(feeds, "b", [])
    sources = write_sources(feeds, ["a", "b"])
    kw = SAMPLE / "keywords_sample.yaml"
    run(make_cfg(sources_path=sources, keywords_path=kw, root=tmp_path))
    new = {
        "t": "Board issues amendments to its climate disclosure standard",
        "u": "https://b.example.com/new",
        "d": "Wed, 14 Oct 2026 09:00:00 +0000",
        "x": "climate disclosure follow-up",
    }
    write_feed(feeds, "b", [new])
    nxt = AS_OF + dt.timedelta(days=1)
    res = run(make_cfg(as_of=nxt, sources_path=sources, keywords_path=kw, root=tmp_path))
    assert "https://b.example.com/new" in all_output_text(res.files)
    assert "同一事件的後續報導" in res.files["daily"].read_text(encoding="utf-8")
    rec = res.day["records"][0]
    assert (
        rec["url"] == "https://b.example.com/new"
        and rec["related_to"]["first_seen"] == "2026-10-14"
    )
    # 同一天重跑，結果不變
    again = run(make_cfg(as_of=nxt, sources_path=sources, keywords_path=kw, root=tmp_path))
    assert again.day["records"] == res.day["records"]


def test_markdown_in_title_is_escaped(make_cfg, tmp_path):
    feeds = tmp_path / "feeds"
    feeds.mkdir()
    item = {
        "t": "碳費 [點我](https://evil.example/x) ## 標題",
        "u": "https://a.example.com/1",
        "d": "Mon, 12 Oct 2026 09:00:00 +0000",
        "x": "碳費新聞",
    }
    write_feed(feeds, "a", [item])
    res = run(make_cfg(sources_path=write_sources(feeds, ["a"]), root=tmp_path))
    md = res.files["daily"].read_text(encoding="utf-8")
    assert "](https://evil.example" not in md
    assert "\\[點我\\]" in md
    assert md_escape("a|b") == "a\\|b"


def test_read_marks_survive_layout_changes(make_cfg):
    res = run(make_cfg())
    path = res.files["tracking"]
    wb = load_workbook(path)
    ws = wb["追蹤表"]
    header = [c.value for c in ws[2]]
    ws.cell(row=3, column=header.index("已讀") + 1).value = "已讀"
    target = ws.cell(row=3, column=1).value
    ws.insert_rows(1)  # 頂端多一列
    ws.insert_rows(6)  # 資料中間插一列空白
    for c in ws[3]:
        if c.value == "已讀":
            c.value = " 已讀 "  # 表頭多空白
    wb.save(path)
    res2 = run(make_cfg())
    rows = {r["文章編號"]: r["已讀"] for r in tracking_rows(res2.files["tracking"])}
    assert rows[target] == "已讀"
    assert "追蹤表：" not in res2.files["daily"].read_text(encoding="utf-8")


@pytest.mark.parametrize("damage", ["corrupt", "no_columns"])
def test_unreadable_tracking_is_backed_up_and_reported(make_cfg, damage):
    res = run(make_cfg())
    path = res.files["tracking"]
    if damage == "corrupt":
        path.write_bytes(b"not a zip")
    else:
        wb = load_workbook(path)
        wb["追蹤表"].delete_cols(1)  # 刪掉「文章編號」欄
        wb.save(path)
    before = path.read_bytes()
    res2 = run(make_cfg())
    backup = path.parent / "情報追蹤表_備份_20261014.xlsx"
    assert backup.read_bytes() == before
    assert "已備份為 情報追蹤表_備份_20261014.xlsx" in res2.files["daily"].read_text(
        encoding="utf-8"
    )
    s = json.loads(res2.files["summary"].read_text(encoding="utf-8"))
    assert any("追蹤表" in x for x in s["data_quality"])
    assert set(res2.files) == {"daily", "weekly", "tracking", "summary"}


@pytest.mark.parametrize(
    "content",
    [
        "id,first_seen\nabc,2026/10/13\n",
        "id,first_seen\nabc,\n",
        "abc,2026-10-13\n",
        "id,first_seen\n,2026-10-13\n",
    ],
)
def test_corrupt_seen_file_is_loud(make_cfg, tmp_path, content):
    seen = tmp_path / "interim" / "seen.csv"
    seen.parent.mkdir(parents=True)
    seen.write_text(content, encoding="utf-8")
    with pytest.raises(IntelDataError, match="seen.csv"):
        run(make_cfg())


def test_seen_file_with_bom_ok(make_cfg, tmp_path):
    seen = tmp_path / "interim" / "seen.csv"
    seen.parent.mkdir(parents=True)
    seen.write_text("id,first_seen\nzzz,2026-10-01\n", encoding="utf-8-sig")
    assert run(make_cfg()).day["records"]


def test_corrupt_day_file_is_loud(make_cfg, tmp_path):
    run(make_cfg())
    (tmp_path / "interim" / "days" / "20261014.json").write_text("{壞掉", encoding="utf-8")
    with pytest.raises(IntelDataError, match="20261014.json"):
        run(make_cfg(as_of=AS_OF + dt.timedelta(days=1)))


def test_cli_refuses_backfill_in_plain_chinese(tmp_path, monkeypatch, capsys):
    import scripts.run_intel as cli

    monkeypatch.setattr(cli, "ROOT", tmp_path)
    out = tmp_path / "data"
    monkeypatch.setattr(cli, "SAMPLE", SAMPLE)
    assert cli.main(["--sample", "--as-of", "2026-10-14"]) == 0
    assert cli.main(["--sample", "--as-of", "2026-10-13"]) == 1
    assert "未執行：" in capsys.readouterr().out
    assert (out / "out" / "intel_sample").exists()
