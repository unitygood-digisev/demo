"""系統 7 端到端測試（用 data/sample/feeds 的假資料，離線執行）。"""

import datetime as dt
import json
import re

import feedparser
from docx import Document
from openpyxl import load_workbook

from intel.evaluate import evaluate, make_sheet, read_labels, to_markdown
from intel.pipeline import run

from .conftest import AS_OF, SAMPLE, set_threshold

URL_RE = re.compile(r"https?://[A-Za-z0-9._~:/?#\[\]@!$&'*+,;=%-]+")


def input_links() -> set[str]:
    return {
        e.link
        for f in ("fake_zh", "fake_en")
        for e in feedparser.parse(SAMPLE / f"{f}.xml").entries
    }


def xlsx_cells(path) -> list[str]:
    wb = load_workbook(path)
    return [
        str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c is not None
    ]


def docx_text(path) -> str:
    d = Document(path)
    parts = [p.text for p in d.paragraphs] + [
        c.text for t in d.tables for r in t.rows for c in r.cells
    ]
    return "\n".join(parts)


def tracking_rows(path) -> list[dict]:
    ws = load_workbook(path)["追蹤表"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1]
    out = []
    for r in rows[2:]:
        if not r or r[0] is None:
            break
        out.append(dict(zip(header, r, strict=False)))
    return out


def all_output_text(files) -> str:
    return "\n".join(
        [
            files["daily"].read_text(encoding="utf-8"),
            docx_text(files["weekly"]),
            "\n".join(xlsx_cells(files["tracking"])),
            files["summary"].read_text(encoding="utf-8"),
        ]
    )


def test_end_to_end_outputs(make_cfg):
    res = run(make_cfg())
    assert set(res.files) == {"daily", "weekly", "tracking", "summary"}
    md = res.files["daily"].read_text(encoding="utf-8")
    assert md.splitlines()[0] == "> 資料等級：公開"
    assert (
        "**今天有 5 則相關動態：法規與準則 3 則、同業動態 1 則、客戶動態 1 則；另有 2 個來源抓取失敗"
        in md
    )
    assert "（假資料）壞掉的來源：RSS 格式壞掉" in md and "（假資料）不存在的來源：找不到檔案" in md
    assert "模擬摘要" in md and "執行編號：intel-20261014-" in md
    # 1002 兩個網址合併成一則、兩個來源都列出
    assert md.count("示範主管機關說明永續報告書查核新制") == 1
    # 2001 與相似標題合併，保留另一個連結
    assert "其他來源連結：https://example.net/fake-wire/9001" in md
    # 低相關放附錄，不丟棄
    assert "附錄：低相關與未列入" in md and "示範主管機關公布年度空品統計" in md
    s = json.loads(res.files["summary"].read_text(encoding="utf-8"))
    assert (
        s["system"] == "intel" and s["classification"] == "公開" and s["metrics"][0]["value"] == 5
    )
    assert s["metrics"][0]["prev_value"] is None and len(s["highlights"]) == 3
    assert len(s["data_quality"]) == 2


def test_every_output_link_comes_from_input(make_cfg):
    res = run(make_cfg())
    found = set(URL_RE.findall(all_output_text(res.files)))
    assert found and found <= input_links(), found - input_links()


def test_model_generated_links_never_reach_output(make_cfg):
    def evil(prompt, text):
        return json.dumps(
            {"摘要": "詳見 https://evil.example/fake 。", "意義": "請看 www.evil.example"}
        )

    res = run(make_cfg(responder=evil))
    text = all_output_text(res.files)
    assert "evil.example" not in text and "已移除模型產生的網址" in text


def test_rerun_same_day_is_identical_and_has_no_duplicates(make_cfg):
    cfg = make_cfg()
    a = run(cfg)
    md1 = a.files["daily"].read_text(encoding="utf-8")
    js1 = a.files["summary"].read_text(encoding="utf-8")
    cells1 = xlsx_cells(a.files["tracking"])
    doc1 = docx_text(a.files["weekly"])
    b = run(make_cfg())
    assert b.run_id == a.run_id
    assert b.files["daily"].read_text(encoding="utf-8") == md1
    assert b.files["summary"].read_text(encoding="utf-8") == js1
    assert xlsx_cells(b.files["tracking"]) == cells1
    assert docx_text(b.files["weekly"]) == doc1
    ids = [r["文章編號"] for r in tracking_rows(b.files["tracking"])]
    assert len(ids) == len(set(ids)) == 8


def test_next_day_same_feeds_has_no_new_items(make_cfg):
    run(make_cfg())
    res = run(make_cfg(as_of=AS_OF + dt.timedelta(days=1)))
    md = res.files["daily"].read_text(encoding="utf-8")
    assert "今天沒有符合的外部情報（本期無符合項目）" in md and "本期無符合項目。" in md
    ids = [r["文章編號"] for r in tracking_rows(res.files["tracking"])]
    assert len(ids) == len(set(ids)) == 8  # 不重複
    s = json.loads(res.files["summary"].read_text(encoding="utf-8"))
    assert s["metrics"][0]["value"] == 5  # 近 7 天仍包含前一天的 5 則
    weekly = docx_text(res.files["weekly"])
    assert "本週有 5 則相關動態" in weekly


def test_read_marks_preserved(make_cfg):
    res = run(make_cfg())
    path = res.files["tracking"]
    wb = load_workbook(path)
    ws = wb["追蹤表"]
    header = [c.value for c in ws[2]]
    col = header.index("已讀") + 1
    target_id = ws.cell(row=3, column=1).value
    assert ws.cell(row=3, column=col).value == "未讀"
    ws.cell(row=3, column=col).value = "已讀"
    wb.save(path)
    res2 = run(make_cfg())
    rows = {r["文章編號"]: r["已讀"] for r in tracking_rows(res2.files["tracking"])}
    assert rows[target_id] == "已讀"
    assert sum(v == "已讀" for v in rows.values()) == 1


def test_llm_only_called_with_public(make_cfg):
    cfg = make_cfg()
    run(cfg)
    assert cfg.llm.levels and set(cfg.llm.levels) == {"公開"}
    assert len(cfg.llm.levels) == 5  # 只對相關的 5 則摘要


def test_relevance_threshold_boundary(make_cfg, config_dir):
    set_threshold(config_dir, "relevance_min_score", 5)  # 碳費、白鷺示範企業 權重剛好 5
    res = run(make_cfg())
    assert res.day["records"] and sum(r["relevant"] for r in res.day["records"]) == 2
    assert "相關門檻：5 分以上" in res.files["daily"].read_text(encoding="utf-8")


def test_max_items_cap(make_cfg, config_dir):
    set_threshold(config_dir, "max_items_per_issue", 2)
    res = run(make_cfg())
    recs = res.day["records"]
    assert sum(r["relevant"] for r in recs) == 2
    assert sum(r["excluded_reason"] == "超過每期上限" for r in recs) == 3
    assert "超過每期上限" in res.files["daily"].read_text(encoding="utf-8")


def test_empty_config_is_loud(make_cfg, tmp_path):
    s = tmp_path / "s.yaml"
    s.write_text("sources: []\n", encoding="utf-8")
    k = tmp_path / "k.yaml"
    k.write_text("keywords: []\n", encoding="utf-8")
    res = run(make_cfg(sources_path=s, keywords_path=k))
    md = res.files["daily"].read_text(encoding="utf-8")
    assert "來源清單是空的" in md and "關鍵字清單是空的" in md and "本期無符合項目" in md


def test_logs_have_no_article_text(make_cfg, tmp_path):
    run(make_cfg())
    raw = (tmp_path / "logs" / "model_calls.jsonl").read_text(encoding="utf-8")
    assert "碳費收費辦法" not in raw and "第三方確信" not in raw


def test_evaluation_math(tmp_path, make_cfg):
    res = run(make_cfg())
    recs = res.day["records"]
    sheet = make_sheet(recs, tmp_path / "標記表.xlsx")
    wb = load_workbook(sheet)
    ws = wb["標記"]
    # 標記：分數 5 的兩則相關、其他都不相關；再加一列亂填
    for row in ws.iter_rows(min_row=2):
        row[5].value = "相關" if row[4].value == 5 else "不相關"
    ws.append(["x", "t", "s", "u", 0, "也許"])
    wb.save(sheet)
    labels, probs = read_labels(sheet)
    assert len(labels) == 8 and len(probs) == 1
    rows, more = evaluate({r["id"]: r["score"] for r in recs}, labels)
    by_t = {r.threshold: r for r in rows}
    assert by_t[1].selected == 5 and by_t[1].recall == 1.0 and by_t[1].precision == 0.4
    assert by_t[5].precision == 1.0 and by_t[5].recall == 1.0
    md = to_markdown(rows, len(labels), probs + more, "2026-10-14")
    assert "| 5 | 2 | 2 | 0 | 0 | 100% | 100% |" in md and "也許" in md
