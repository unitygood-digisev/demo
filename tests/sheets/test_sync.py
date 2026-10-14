"""階段 4b：同步到 Google 試算表（全部用測試替身，不連網路）。"""

import datetime as dt
import json
import os
import shutil
import socket
import subprocess
from pathlib import Path

import pytest
import yaml

from digest.pipeline import RunConfig as DigestConfig
from digest.pipeline import run as run_digest
from intel.pipeline import RunConfig as IntelConfig
from intel.pipeline import run as run_intel
from progress.pipeline import RunConfig as ProgressConfig
from progress.pipeline import run as run_progress
from shared.config import load_dotenv
from sheets.client import FakeClient, GspreadClient, _pad
from sheets.content import SyncError, merge_history, read_sheet_table
from sheets.sync import SyncConfig, sync

ROOT = Path(__file__).resolve().parents[2]
AS_OF = dt.date(2026, 10, 14)
SAMPLE_P = ROOT / "data" / "sample" / "progress"
SAMPLE_F = ROOT / "data" / "sample" / "feeds"
TITLES = ["週報", "進度預警", "進度明細", "外部情報", "說明", "指標歷史"]


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("測試中不允許連網路")

    monkeypatch.setattr(socket, "create_connection", boom)
    monkeypatch.setattr(socket.socket, "connect", boom)


@pytest.fixture(scope="module")
def outputs(tmp_path_factory):
    """用假資料實際跑三個系統，產生本機報表。"""
    t = tmp_path_factory.mktemp("w")
    for d in (dt.date(2026, 9, 30), dt.date(2026, 10, 7), AS_OF):
        run_progress(
            ProgressConfig(
                as_of=d,
                input_path=SAMPLE_P / f"項目進度看板_假資料_{d.strftime('%Y%m%d')}.xlsx",
                interim_dir=t / "pi",
                out_dir=t / "progress",
                log_dir=t / "logs",
            )
        )
    run_intel(
        IntelConfig(
            as_of=AS_OF,
            sources_path=SAMPLE_F / "sources_sample.yaml",
            keywords_path=SAMPLE_F / "keywords_sample.yaml",
            interim_dir=t / "ii",
            out_dir=t / "intel",
            offline=True,
            log_dir=t / "logs",
        )
    )
    run_digest(
        DigestConfig(
            as_of=AS_OF,
            out_dir=t / "digest",
            log_dir=t / "logs",
            summary_overrides={
                "progress": t / "progress" / "summary.json",
                "intel": t / "intel" / "summary.json",
            },
        )
    )
    return t


@pytest.fixture
def work(outputs, tmp_path):
    """每個測試一份可以亂改的報表副本。"""
    w = tmp_path / "w"
    for sub in ("digest", "progress", "intel"):
        shutil.copytree(outputs / sub, w / sub)
    return w


@pytest.fixture
def cfg(work, tmp_path):
    c = tmp_path / "config"
    shutil.copytree(ROOT / "config", c)

    def make(client, **kw):
        base = dict(
            as_of=AS_OF,
            config_dir=c,
            digest_dir=work / "digest",
            progress_dir=work / "progress",
            intel_dir=work / "intel",
            client=client,
            env_file=tmp_path / "no.env",
        )
        base.update(kw)
        return SyncConfig(**base)

    make.config_dir = c
    return make


def _s(rows):
    return [[str(v) for v in r] for r in rows]


def digest_json(work):
    return json.loads(
        (work / "digest" / "管理層週報_資料_20261014.json").read_text(encoding="utf-8")
    )


# ---- 內容等於本機報表 ----
def test_every_tab_equals_local_reports(cfg, work):
    fake = FakeClient()
    res = sync(cfg(fake))
    assert res.written == TITLES and res.notes == []
    d = digest_json(work)
    for t in TITLES[:-1]:
        assert (
            fake.tabs[t][0][0] == f"資料等級：內部一般｜資料截至 2026-10-14｜執行編號 {d['run_id']}"
        )
    tab = _s(fake.tabs["週報"])
    for i, s in enumerate(d["sentences"], 1):
        assert [f"{i}. {s}"] in tab
    for m in d["metrics"]:
        assert [m["label"], str(m["this"]), str(m["prev"]), m["change"], m["unit"] or "—"] in tab
    for x in d["decisions"] + d["data_quality"]:
        assert [x] in tab
    for i in d["intel"]:
        assert [i["title"], i["url"]] in tab
    xlsx = work / "progress" / "進度看板_20261014.xlsx"
    for title, sheet in (("進度預警", "預警清單"), ("進度明細", "項目明細")):
        header, rows = read_sheet_table(xlsx, sheet)
        assert fake.tabs[title][2] == header and _s(fake.tabs[title][3:]) == _s(rows)
    header, rows = read_sheet_table(work / "intel" / "情報追蹤表.xlsx", "追蹤表")
    keep = [
        "首次看到",
        "來源",
        "標題",
        "連結",
        "分類",
        "是否相關",
        "摘要",
        "對優樂地的意義",
        "分數",
    ]
    assert fake.tabs["外部情報"][2] == keep
    assert _s(fake.tabs["外部情報"][3:]) == _s([[r[header.index(c)] for c in keep] for r in rows])
    assert [r[0] for r in fake.tabs["說明"][3:]] == d["footer"]


def test_excluded_fields_never_synced(cfg):
    fake = FakeClient()
    sync(cfg(fake))
    text = " ".join(str(c) for rows in fake.tabs.values() for r in rows for c in r)
    for word in ("不續約", "備註內容", "交初稿", "邀講者", "已讀", "未讀"):
        assert word not in text


# ---- 指標歷史 ----
def test_history_full_flow_and_numbers(cfg, work):
    fake = FakeClient()
    sync(cfg(fake))
    sync(cfg(fake))  # 同一週重跑不重複
    rows = fake.tabs["指標歷史"]
    assert len(rows) == 2 and rows[1][0] == "2026-10-14"
    assert isinstance(rows[1][1], int)  # 數字以數字寫入，才能畫趨勢
    # 下一週：改週報資料的日期模擬下週
    d = digest_json(work)
    d["as_of"] = "2026-10-21"
    (work / "digest" / "管理層週報_資料_20261021.json").write_text(
        json.dumps(d, ensure_ascii=False), encoding="utf-8"
    )
    sync(cfg(fake, as_of=dt.date(2026, 10, 21)))
    assert [r[0] for r in fake.tabs["指標歷史"][1:]] == ["2026-10-14", "2026-10-21"]


def _apply(existing, cells):
    fake = FakeClient({"x": existing})
    fake.update_cells("x", cells)
    return fake.tabs["x"]


def test_history_keeps_manager_content():
    h = ["週報日期", "紅", "資料等級"]
    existing = [
        ["週報日期", "紅", "資料等級", "主管備註欄"],
        ["2026/10/07", 2, "內部一般", "第一週"],
        [],
        ["", "這是主管的註記", "-5% 待確認", "007"],
        ["2026-10-14", "=B2+1", "內部一般", "公式要保留", "表頭外的格子"],
    ]
    cells, notes = merge_history(existing, h, ["2026-10-14", 3, "內部一般"])
    assert notes == []
    # 只改本週那一列的系統欄位，其他格子（含主管的公式、文字、註記列）一個都不碰
    assert {(r, c) for r, c, _ in cells} == {(4, 0), (4, 1), (4, 2)}
    out = _apply(existing, cells)
    assert out[4] == ["2026-10-14", 3, "內部一般", "公式要保留", "表頭外的格子"]
    assert out[:4] == existing[:4]
    # 新的一週接在最後一列之後；2026/10/07 視為已有的那一週
    cells, _ = merge_history(out, h, ["2026-10-21", 4, "內部一般"])
    assert {r for r, _, _ in cells} == {5}
    cells, _ = merge_history(out, h, ["2026-10-07", 9, "內部一般"])
    assert {r for r, _, _ in cells} == {1}


@pytest.mark.parametrize("first", [["日期", "紅"], ["", "", ""]])
def test_history_unknown_layout_not_touched(cfg, first):
    # 第 1 格空白但分頁有內容，也不是系統寫的：停下來，不覆蓋
    fake = FakeClient({"指標歷史": [first, ["2026-10-07", 1]]})
    with pytest.raises(SyncError, match="不是系統寫的"):
        sync(cfg(fake))
    assert fake.tabs["指標歷史"] == [first, ["2026-10-07", 1]]


def test_history_column_renamed_keeps_both(cfg):
    h = ["週報日期", "紅色預警項目數", "資料等級"]
    existing = [["週報日期", "紅色預警（舊名）", "資料等級"], ["2026-10-07", 2, "內部一般"]]
    cells, _ = merge_history(existing, h, ["2026-10-14", 3, "內部一般"])
    out = _apply(existing, cells)
    assert out[0] == ["週報日期", "紅色預警（舊名）", "資料等級", "紅色預警項目數"]
    assert out[1][1] == 2 and out[2][3] == 3


# ---- 不覆蓋使用者的分頁 ----
def test_user_tabs_untouched(cfg):
    fake = FakeClient({"主管筆記": [["我的筆記"]]})
    sync(cfg(fake))
    assert fake.tabs["主管筆記"] == [["我的筆記"]] and "主管筆記" not in fake.writes


def test_same_name_user_tab_stops_before_writing(cfg):
    fake = FakeClient({"說明": [["這是主管自己寫的說明"]]})
    with pytest.raises(SyncError, match="不是系統寫的"):
        sync(cfg(fake))
    assert fake.writes == []


# ---- 設定與 .env ----
def test_dotenv_loaded_and_paths_cleaned(cfg, tmp_path, monkeypatch):
    monkeypatch.delenv("MGMT_SHEET_ID", raising=False)
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    monkeypatch.setenv("HOME", str(tmp_path))
    key = tmp_path / "secret dir" / "key.json"
    key.parent.mkdir()
    key.write_text("{}", encoding="utf-8")
    env = tmp_path / ".env"
    env.write_text(
        '# 註解\nMGMT_SHEET_ID="abc123"\nGOOGLE_APPLICATION_CREDENTIALS="~/secret dir/key.json"\n',
        encoding="utf-8",
    )
    seen = {}

    class Spy:
        def __init__(self, path, sid):
            seen.update(path=path, sid=sid)
            raise SyncError("停在這裡")

    monkeypatch.setattr("sheets.sync.GspreadClient", Spy)
    with pytest.raises(SyncError, match="停在這裡"):
        sync(cfg(None, env_file=env))
    assert seen == {"path": key, "sid": "abc123"}


def test_dotenv_does_not_override_existing(tmp_path, monkeypatch):
    monkeypatch.setenv("MGMT_SHEET_ID", "from-env")
    env = tmp_path / ".env"
    env.write_text("MGMT_SHEET_ID=from-file\n", encoding="utf-8")
    load_dotenv(env)
    assert os.environ["MGMT_SHEET_ID"] == "from-env"


def test_not_configured_is_skipped(cfg, monkeypatch):
    monkeypatch.delenv("MGMT_SHEET_ID", raising=False)
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    with pytest.raises(SyncError, match="略過同步"):
        sync(cfg(None))


@pytest.mark.parametrize(
    "mutate, msg",
    [
        (lambda c: c.update(enabled=False), "enabled 為 false"),
        (lambda c: c.update(enabled="no"), "只能寫 true 或 false"),
        (lambda c: c["tabs"].pop("history"), "tabs 缺少：history"),
        (lambda c: c["tabs"].update(alerts="週報"), "不能重複"),
        (lambda c: c["tabs"].update(info="../raw/x"), "不合法"),
        (lambda c: c["tabs"].update(info="a/b"), "不合法"),
        (lambda c: c["tabs"].update(info="長" * 101), "不合法"),
        (lambda c: c["tabs"].update(info=123), "不合法"),
    ],
)
def test_bad_config(cfg, mutate, msg):
    p = cfg.config_dir / "sheets.yaml"
    c = yaml.safe_load(p.read_text(encoding="utf-8"))
    mutate(c)
    p.write_text(yaml.safe_dump(c, allow_unicode=True), encoding="utf-8")
    with pytest.raises(SyncError, match=msg):
        sync(cfg(FakeClient()))


def test_config_not_a_mapping(cfg):
    (cfg.config_dir / "sheets.yaml").write_text("- a\n- b\n", encoding="utf-8")
    with pytest.raises(SyncError, match="格式錯誤"):
        sync(cfg(FakeClient()))


# ---- 本機報表壞掉 ----
@pytest.mark.parametrize("content", ['{"as_of": "2026-10-14"}', "[1]", b"\xff\xfe\x00bad", '{"壞'])
def test_broken_digest_json_is_plain(cfg, work, content):
    p = work / "digest" / "管理層週報_資料_20261014.json"
    p.write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
    with pytest.raises(SyncError, match="週報資料檔"):
        sync(cfg(FakeClient()))


@pytest.mark.parametrize("content", [b"", b"not a zip"])
def test_broken_progress_xlsx_is_plain(cfg, work, content):
    (work / "progress" / "進度看板_20261014.xlsx").write_bytes(content)
    with pytest.raises(SyncError, match="已損壞或不是 Excel 檔"):
        sync(cfg(FakeClient()))


def test_intel_required_columns(cfg, work):
    from openpyxl import load_workbook

    p = work / "intel" / "情報追蹤表.xlsx"
    wb = load_workbook(p)
    ws = wb["追蹤表"]
    for c in ws[2]:
        if c.value == "標題":
            c.value = "新聞標題"
    wb.save(p)
    with pytest.raises(SyncError, match="缺少欄位：標題"):
        sync(cfg(FakeClient()))


def test_stale_system_not_written_as_this_week(cfg, work):
    d = digest_json(work)
    d["systems"]["intel"] = {
        "name": "外部情報摘要",
        "status": "failed",
        "text": "本週執行失敗，未納入",
    }
    (work / "digest" / "管理層週報_資料_20261014.json").write_text(
        json.dumps(d, ensure_ascii=False), encoding="utf-8"
    )
    fake = FakeClient()
    res = sync(cfg(fake))
    assert "本週執行失敗" in fake.tabs["外部情報"][1][0] and len(fake.tabs["外部情報"]) == 2
    assert any("外部情報" in n for n in res.notes)


def test_missing_digest_for_date(cfg):
    with pytest.raises(SyncError, match="請先產生週報"):
        sync(cfg(FakeClient(), as_of=AS_OF + dt.timedelta(days=1)))


def test_missing_progress_report_noted(cfg, tmp_path):
    fake = FakeClient()
    res = sync(cfg(fake, progress_dir=tmp_path / "none"))
    assert "未更新" in fake.tabs["進度預警"][1][0] and res.notes


# ---- 寫入失敗 ----
def test_partial_write_failure_lists_tabs(cfg):
    class Broken(FakeClient):
        def write(self, title, rows):
            if title == "進度明細":
                raise ConnectionError("x")
            super().write(title, rows)

    with pytest.raises(
        SyncError,
        match="部分已更新：已寫入 週報、進度預警；未寫入 進度明細、外部情報、說明、指標歷史",
    ):
        sync(cfg(Broken()))


def test_gspread_write_overwrites_without_clearing():
    calls = {}

    class WS:
        row_count, col_count = 100, 12

        def get_all_values(self, **kw):
            return [["舊"] * 4] * 5

        def clear(self):
            raise AssertionError("不應先清空")

        def resize(self, **kw):
            calls["resize"] = kw

        def update(self, values, range_name, value_input_option):
            calls.update(values=values, opt=value_input_option)

    from gspread.utils import ValueInputOption

    c = object.__new__(GspreadClient)
    c._raw = ValueInputOption.raw
    c._ws = lambda title: WS()
    c.write("x", [["=HYPERLINK(1)", 5, None], ["a"]])
    assert calls["opt"] == ValueInputOption.raw
    assert calls["values"][0] == ["=HYPERLINK(1)", 5, "", ""] and len(calls["values"]) == 5
    assert calls["values"][4] == ["", "", "", ""]  # 舊內容多出的列被空白蓋掉


def test_gspread_update_cells_raw_only_given_cells():
    calls = {}

    class WS:
        row_count, col_count = 3, 3

        def resize(self, **kw):
            calls["resize"] = kw

        def batch_update(self, data, value_input_option):
            calls.update(data=data, opt=value_input_option)

    from gspread.utils import ValueInputOption

    c = object.__new__(GspreadClient)
    c._raw = ValueInputOption.raw
    c._ws = lambda title: WS()
    c.update_cells("x", [(4, 0, "2026-10-14"), (4, 3, "=IMPORTXML(1)"), (4, 1, 5)])
    assert calls["opt"] == ValueInputOption.raw  # 原樣：日期是文字、= 開頭不會變公式
    assert calls["data"] == [
        {"range": "A5", "values": [["2026-10-14"]]},
        {"range": "D5", "values": [["=IMPORTXML(1)"]]},
        {"range": "B5", "values": [[5]]},
    ]
    assert calls["resize"] == {"rows": 5, "cols": 4}


def test_pad_grows_to_cover_old():
    assert _pad([["a", 1]], 3, 3) == [["a", 1, ""], ["", "", ""], ["", "", ""]]


# ---- 倉庫衛生與一鍵執行 ----
def test_no_credentials_in_repo():
    files = subprocess.run(
        ["git", "ls-files"], cwd=ROOT, capture_output=True, text=True, check=True
    ).stdout.split()
    needles = [
        "BEGIN " + "PRIVATE KEY",
        '"private' + '_key"',
        "docs.google.com/spreadsheets" + "/d/1",
    ]
    for f in files:
        p = ROOT / f
        if p.suffix in (".json", ".py", ".yaml", ".md", ".txt", ".toml", ".example", ".csv"):
            text = p.read_text(encoding="utf-8", errors="ignore")
            for n in needles:
                assert n not in text, f


def test_env_and_preview_are_gitignored():
    for path in (".env", "data/out/sheets_preview/週報.csv"):
        r = subprocess.run(["git", "check-ignore", path], cwd=ROOT, capture_output=True, text=True)
        assert r.returncode == 0, path


def test_preview_cli_offline_and_safe(monkeypatch, outputs, tmp_path, capsys):
    import scripts.sync_sheets as cli

    monkeypatch.setattr(cli, "ROOT", tmp_path)
    monkeypatch.setattr("sheets.sync.OUT", outputs)
    monkeypatch.setattr(
        cli,
        "SyncConfig",
        lambda as_of: SyncConfig(
            as_of=as_of,
            digest_dir=outputs / "digest",
            progress_dir=outputs / "progress",
            intel_dir=outputs / "intel",
        ),
    )
    assert cli.main(["--as-of", "2026-10-14", "--preview"]) == 0
    files = sorted(p.name for p in (tmp_path / "data" / "out" / "sheets_preview").glob("*.csv"))
    assert files == sorted(f"{t}.csv" for t in TITLES)
    assert cli.main(["--as-of", "2026-13-01"]) == 1
    assert "看不懂" in capsys.readouterr().out
    assert cli._safe("=1+1") == "'=1+1" and cli._safe("@x") == "'@x" and cli._safe("正常") == "正常"


def test_run_weekly_messages(monkeypatch, capsys):
    import scripts.run_weekly as weekly
    from scripts import sync_sheets

    monkeypatch.setattr(weekly, "STEPS", [])
    monkeypatch.setattr(weekly, "DIGEST", ("管理層週報", lambda a: 0))
    monkeypatch.setattr(weekly, "SHEETS", ("同步 Google 試算表", lambda a: sync_sheets.SKIPPED))
    assert weekly.main([]) == 0
    assert "全部完成。（Google 試算表尚未設定或已停用，本週沒有同步）" in capsys.readouterr().out
    monkeypatch.setattr(weekly, "SHEETS", ("同步 Google 試算表", lambda a: 1))
    assert weekly.main([]) == 1
    monkeypatch.setattr(weekly, "DIGEST", ("管理層週報", lambda a: 1))
    assert weekly.main([]) == 1
    assert "週報沒有產生，所以沒有同步 Google 試算表" in capsys.readouterr().out
    assert weekly.main(["--as-of", "2026-13-01"]) == 1


# ---- 第二輪驗收追加 ----
def _serial(d: dt.date) -> int:
    return (d - dt.date(1899, 12, 30)).days


@pytest.mark.parametrize(
    "old", [_serial(dt.date(2026, 10, 14)), "46309", "'2026-10-14", "2026/10/14"]
)
def test_history_date_serial_not_duplicated(cfg, old):
    # Google 把日期存成日期格式時，讀回來可能是序號；同一週不能多出一列
    assert _serial(dt.date(2026, 10, 14)) == 46309
    fake = FakeClient({"指標歷史": [["週報日期", "紅色預警項目數"], [old, 1]]})
    sync(cfg(fake))
    sync(cfg(fake))
    rows = fake.tabs["指標歷史"]
    assert [r[0] for r in rows[1:] if r and r[0]] == ["2026-10-14"]


def test_history_first_write_is_cells_too():
    # 分頁還是空的：表頭＋本週一列，同樣用「原樣」寫格子
    cells, _ = merge_history([], ["週報日期", "備註"], ["2026-10-14", "=IMPORTXML(1)"])
    assert _apply([], cells) == [["週報日期", "備註"], ["2026-10-14", "=IMPORTXML(1)"]]


@pytest.mark.parametrize("v", ["2026年10月14日", "2026-10-14 00:00:00", "2026/10/14"])
def test_history_other_date_displays_same_week(v):
    cells, _ = merge_history([["週報日期", "紅"], [v, 1]], ["週報日期", "紅"], ["2026-10-14", 2])
    assert {r for r, _, _ in cells} == {1}


def test_history_ambiguous_date_not_guessed():
    # 10/14/2026 月日順序不確定，不猜：保留原列，本週另寫一列
    cells, _ = merge_history(
        [["週報日期", "紅"], ["10/14/2026", 1]], ["週報日期", "紅"], ["2026-10-14", 2]
    )
    assert {r for r, _, _ in cells} == {2}


def test_gspread_read_formulas_asks_for_text_dates():
    seen = {}

    class WS:
        def get_all_values(self, **kw):
            seen.update(kw)
            return []

    from gspread.utils import DateTimeOption, ValueRenderOption

    c = object.__new__(GspreadClient)
    c._formula, c._date_text = ValueRenderOption.formula, DateTimeOption.formatted_string
    c._ws = lambda title: WS()
    c.read("x", formulas=True)
    assert seen["date_time_render_option"] == DateTimeOption.formatted_string


def test_dotenv_inline_comment_and_empty_env(tmp_path, monkeypatch):
    monkeypatch.setenv("MGMT_SHEET_ID", "")  # 空白的環境變數要被 .env 補上
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    monkeypatch.delenv("X_QUOTED", raising=False)
    env = tmp_path / ".env"
    env.write_text(
        "MGMT_SHEET_ID=abc123   # 試算表 ID\n"
        "GOOGLE_APPLICATION_CREDENTIALS=C:/keys/k#1.json\n"
        'X_QUOTED="a # b"  # 註解\n',
        encoding="utf-8",
    )
    load_dotenv(env)
    assert os.environ["MGMT_SHEET_ID"] == "abc123"
    assert os.environ["GOOGLE_APPLICATION_CREDENTIALS"] == "C:/keys/k#1.json"
    assert os.environ["X_QUOTED"] == "a # b"


def test_dotenv_not_utf8_is_plain(cfg, tmp_path, monkeypatch):
    monkeypatch.delenv("MGMT_SHEET_ID", raising=False)
    env = tmp_path / ".env"
    env.write_bytes("MGMT_SHEET_ID=試算表\n".encode("big5"))
    with pytest.raises(SyncError, match="UTF-8"):
        sync(cfg(None, env_file=env))


def test_digest_date_must_match_filename(cfg, work):
    d = digest_json(work)
    d["as_of"] = "2026-10-07"
    (work / "digest" / "管理層週報_資料_20261014.json").write_text(
        json.dumps(d, ensure_ascii=False), encoding="utf-8"
    )
    with pytest.raises(SyncError, match="和檔名不符"):
        sync(cfg(FakeClient()))


def test_progress_header_shift_is_plain(cfg, work):
    from openpyxl import load_workbook

    p = work / "progress" / "進度看板_20261014.xlsx"
    wb = load_workbook(p)
    wb["項目明細"].insert_rows(1)  # 表頭被往下推一列
    wb.save(p)
    fake = FakeClient()
    with pytest.raises(SyncError, match="缺少欄位"):
        sync(cfg(fake))
    assert fake.writes == []


def test_dotenv_last_value_wins_and_plain_errors(tmp_path, monkeypatch):
    monkeypatch.delenv("X_DUP", raising=False)
    env = tmp_path / ".env"
    env.write_text("X_DUP=first\nX_DUP=second\n", encoding="utf-8")
    load_dotenv(env)
    assert os.environ["X_DUP"] == "second"
    env.write_text('X_BAD="abc\n', encoding="utf-8")
    with pytest.raises(ValueError, match="第 1 行的引號沒有成對"):
        load_dotenv(env)
    env.write_text("X_NUL=a\x00b\n", encoding="utf-8")
    with pytest.raises(ValueError, match="特殊字元"):
        load_dotenv(env)


@pytest.mark.parametrize("bad", [None, [], {"status": None}, "ok"])
def test_digest_systems_items_checked(cfg, work, bad):
    d = digest_json(work)
    d["systems"]["progress"] = bad
    (work / "digest" / "管理層週報_資料_20261014.json").write_text(
        json.dumps(d, ensure_ascii=False), encoding="utf-8"
    )
    with pytest.raises(SyncError, match="systems"):
        sync(cfg(FakeClient()))
