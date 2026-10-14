"""產生系統 1 的假資料：與公司現有「項目進度看板」相同格式的三週版本。

所有案子、客戶、人名都是虛構的，並標示「假資料」。
- 項目進度看板_假資料_20260930.xlsx
- 項目進度看板_假資料_20261007.xlsx
- 項目進度看板_假資料_20261014.xlsx（主要測試版本，基準日 2026-10-14）
刻意放入的錯誤與預期觸發的規則，見 data/sample/README.md。

執行：.venv/bin/python scripts/make_sample_data.py
"""

from __future__ import annotations

import copy
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "data" / "sample" / "progress"
D = dt.date
HEADER = [
    "客戶／案子",
    "類別",
    "階段",
    "負責人",
    "近期下一步",
    "下一個日期",
    "卡點／風險",
    "備註",
    "距今天數",
]
STAGES = "規劃中,洽談中,進行中,卡關,結案中,已完成,暫緩"
SECRET = "（假資料）客戶內部考量，明年可能不續約"  # 用來測試「卡點／風險」不會出現在輸出


def r(name, cat, stage, owner, step, date, risk=None, note=None):
    return [name, cat, stage, owner, step, date, risk, note]


# 2026-10-14 版本（主要測試資料）
CURRENT = [
    r(
        "（假資料）甲案",
        "永續報告書",
        "進行中",
        "顧問甲",
        "（假資料）交初稿",
        D(2026, 10, 13),
        SECRET,
    ),  # P01 邊界：過期 1 天
    r(
        "（假資料）乙案", "教育訓練", "進行中", "顧問甲", "（假資料）上課", D(2026, 10, 14)
    ),  # P02：今天到期
    r(
        "（假資料）丙案", "結案點交", "結案中", "顧問乙", "（假資料）點交", D(2026, 10, 21)
    ),  # P02 邊界：剛好 7 天
    r(
        "（假資料）丁案", "工作坊", "進行中", "顧問乙", "（假資料）工作坊", D(2026, 10, 22)
    ),  # 8 天：不觸發
    r(
        "（假資料）戊案",
        "系統開發",
        "卡關",
        "顧問丙",
        "（假資料）等對方回覆",
        D(2026, 10, 30),
        SECRET,
    ),  # P03
    r("（假資料）己案", "SROI", "進行中", "顧問甲（與顧問乙）", "（假資料）訪談", None),  # 缺日期
    r(
        "（假資料）庚案", "活動", "已完成", "顧問丙", "（假資料）已結束", D(2026, 10, 1)
    ),  # 已完成：不觸發
    r(
        "（假資料）辛案", "BD", "洽談中", "顧問丁", "（假資料）報價", D(2026, 10, 10)
    ),  # P01：洽談中也算
    r(
        "（假資料）壬案", "碳盤查", "進行中", "顧問甲", "（假資料）收資料", D(2026, 11, 5)
    ),  # P04：自 9/30 未變（14 天）
    r(
        "（假資料）癸案", "標案", "進行中", "顧問甲", "（假資料）備標", D(2026, 11, 10)
    ),  # 自 10/7 未變（7 天）：不觸發
    r("（假資料）子案", "獎項", "進行中", "顧問甲／顧問丁", "（假資料）送件", D(2026, 11, 20)),
    r("（假資料）丑案", "行政", "暫緩", "行政", "（假資料）暫停", None),  # 暫緩：不需日期
    r(
        "（假資料）寅案",
        "論壇",
        "進行中",
        "顧問甲",
        "（假資料）邀講者",
        D(2026, 12, 1),
        None,
        "（假資料）備註內容",
    ),  # 顧問甲共 7 案 → P05
    r("（假資料）卯案", "教育訓練", "進行中", "顧問乙", "（假資料）排課", D(2026, 11, 15)),
    r("（假資料）辰案", "結案點交", "結案中", "顧問乙", "（假資料）結報", D(2026, 11, 16)),
    r(
        "（假資料）巳案", "工作坊", "進行中", "顧問乙", "（假資料）場勘", D(2026, 11, 17)
    ),  # 顧問乙剛好 5 案：不觸發
    r(
        "（假資料）午案", "活動", "進行中", "全體專案人員", "（假資料）彩排", D(2026, 11, 18)
    ),  # 非人名：不計
    # ---- 刻意放入的資料錯誤 ----
    r("（假資料）未案", "標案", "進行中", "顧問丁", "（假資料）送件", "2026/13/01"),  # 日期錯誤
    r(
        "（假資料）甲案", "合作", "進行中", "顧問丁", "（假資料）重複的案名", D(2026, 11, 20)
    ),  # 案名重複
    r(
        "（假資料）酉案", "專案", "完成了", "顧問丁", "（假資料）不在選單的階段", D(2026, 11, 21)
    ),  # 階段不在選單
    r(None, "詢價", "規劃中", "顧問丁", "（假資料）案名空白", None),  # 缺案名
    r(
        "（假資料）戌案", "補助", "進行中", None, "（假資料）沒有負責人", D(2026, 11, 30)
    ),  # 缺負責人
]


# 每週都有更新「近期下一步」的案子（壬案、癸案除外），避免假資料裡每案都「長期沒有變化」
STEADY = {"（假資料）壬案", "（假資料）癸案"}


def _touch(rows: list[list], tag: str) -> None:
    for x in rows:
        if x[0] not in STEADY:
            x[4] = f"{x[4]}（{tag}）"


def version_1007() -> list[list]:
    rows = copy.deepcopy(CURRENT[:17])
    _touch(rows, "10/7 版")
    by = {x[0]: x for x in rows}
    by["（假資料）甲案"][5] = D(2026, 10, 20)  # 10/7 時尚未到期 → 10/14 新增 P01
    by["（假資料）庚案"][2] = "進行中"  # 10/7 時逾期 → 10/14 已完成，預警解除
    by["（假資料）庚案"][5] = D(2026, 10, 6)
    rows.append(
        r("（假資料）亥案", "詢價", "洽談中", "顧問丁", "（假資料）等回覆", D(2026, 10, 30))
    )  # 10/14 消失
    return rows


def version_0930() -> list[list]:
    rows = version_1007()
    for x in rows:
        if x[0] not in STEADY:
            x[4] = x[4].replace("（10/7 版）", "（9/30 版）")
    by = {x[0]: x for x in rows}
    by["（假資料）癸案"][5] = D(2026, 10, 31)  # 9/30 與之後不同 → 癸案只維持 7 天
    return rows


def write(rows: list[list], board_date: dt.date) -> Path:
    wb = Workbook()
    ov = wb.active
    ov.title = "總覽"
    ov.append(["項目看板總覽（假資料）"])
    ws = wb.create_sheet("項目看板")
    ws.append([f"項目看板：（假資料）{board_date.month}/{board_date.day} 專案組會議紀要"])
    ws.append(
        [
            f"來源：{board_date.year}/{board_date.month}/{board_date.day} 會議紀錄（假資料）。黃底可修改；階段為下拉選"
        ]
    )
    ws.append([])
    ws.append(HEADER)
    for i, row in enumerate(rows, start=5):
        ws.append([*row, f'=IF(OR(F{i}="",C{i}="已完成"),"",F{i}-TODAY())'])
        if isinstance(row[5], dt.date):
            ws.cell(row=i, column=6).number_format = "yyyy-mm-dd"
    dv = DataValidation(type="list", formula1=f'"{STAGES}"', allow_blank=True)
    dv.add(f"C5:C{4 + len(rows)}")
    ws.add_data_validation(dv)
    todo = wb.create_sheet("待辦事項")
    todo.append(["我的待辦事項（假資料，系統不讀）"])
    wb.properties.title = "假資料"
    wb.properties.creator = "unity-mgmt-suite"
    wb.properties.created = dt.datetime(2000, 1, 1)
    wb.properties.modified = dt.datetime(2000, 1, 1)
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"項目進度看板_假資料_{board_date.strftime('%Y%m%d')}.xlsx"
    wb.save(path)
    return path


def main() -> None:
    for rows, day in (
        (version_0930(), D(2026, 9, 30)),
        (version_1007(), D(2026, 10, 7)),
        (CURRENT, D(2026, 10, 14)),
    ):
        print(f"已產生 {write(rows, day)}")


if __name__ == "__main__":
    main()
